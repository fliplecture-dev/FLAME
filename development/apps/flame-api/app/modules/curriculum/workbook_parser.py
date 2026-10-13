"""Bounded, read-only XLSX parsing for curriculum import previews."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

HEADERS = (
    "Curriculum Code", "Level", "Subject", "Domain", "Category", "Topic",
    "Subtopic", "Learning Outcome", "Source URL",
)
OPTIONAL_OUTCOME_HEADERS = HEADERS[:7] + HEADERS[8:]
REQUIRED = (0, 1, 2, 3, 5, 6)
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_SHEETS = 32
MAX_ROWS = 20_000
MAX_SHEET_ROWS = 25_000
MAX_ISSUES = 200
ALL_LEVELS = "__ALL_LEVELS__"
MAX_FIELD_LENGTHS = (80, 80, 255, 255, 255, 255, 255, 2000, 2000)


class WorkbookRejected(ValueError):
    """A malformed or oversized workbook cannot be safely parsed."""

    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class RowIssue:
    sheet: str
    row: int
    column: str | None
    code: str


@dataclass(frozen=True)
class ImportRow:
    sheet: str
    row: int
    values: tuple[str | None, ...]
    path_key: tuple[str | None, ...]
    duplicate_of: tuple[str, int] | None


@dataclass(frozen=True)
class ParsedWorkbook:
    file_hash: str
    rows: tuple[ImportRow, ...]
    issues: tuple[RowIssue, ...]
    issue_count: int
    duplicate_count: int
    sheet_counts: dict[str, int]

    @property
    def valid(self) -> bool:
        return self.issue_count == 0


def normalize_text(value: str) -> str:
    return " ".join(unicodedata.normalize("NFC", value).split())


def match_key(value: str | None) -> str | None:
    return normalize_text(value).casefold() if value else None


def _check_zip(payload: bytes) -> None:
    if len(payload) > MAX_FILE_BYTES:
        raise WorkbookRejected("file_too_large")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_ENTRIES:
                raise WorkbookRejected("too_many_zip_entries")
            if len({member.filename for member in members}) != len(members):
                raise WorkbookRejected("duplicate_zip_entry")
            if sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
                raise WorkbookRejected("expanded_file_too_large")
            for member in members:
                name = member.filename.replace("\\", "/")
                path = PurePosixPath(name)
                if name.startswith("/") or ".." in path.parts or member.flag_bits & 1:
                    raise WorkbookRejected("unsafe_zip_entry")
                lower = name.casefold()
                if (lower.startswith("xl/externallinks/") or lower.startswith("xl/embeddings/")
                    or lower.startswith("xl/activex/") or lower.endswith("vbaproject.bin")):
                    raise WorkbookRejected("unsupported_workbook_content")
            if "xl/workbook.xml" not in archive.namelist():
                raise WorkbookRejected("invalid_workbook")
    except BadZipFile as exc:
        raise WorkbookRejected("invalid_workbook") from exc


def parse_workbook(
    payload: bytes,
    *,
    filename: str,
    sheet_levels: dict[str, str],
    expected_code: str,
) -> ParsedWorkbook:
    """Parse selected sheets; return bounded row issues without writing catalogue data.

    ``sheet_levels`` maps selected sheet names to expected Level display names.
    The target curriculum ID/code lookup is the caller's responsibility.
    """
    if not filename.casefold().endswith(".xlsx"):
        raise WorkbookRejected("unsupported_file_type")
    if not 1 <= len(sheet_levels) <= MAX_SHEETS or any(
        not name or not level or (level != ALL_LEVELS and len(normalize_text(level)) > 80)
        for name, level in sheet_levels.items()
    ):
        raise WorkbookRejected("invalid_sheet_mapping")
    if not expected_code or len(normalize_text(expected_code)) > 80:
        raise WorkbookRejected("invalid_curriculum_code")
    _check_zip(payload)
    try:
        workbook = load_workbook(BytesIO(payload), read_only=False, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, TypeError) as exc:
        raise WorkbookRejected("invalid_workbook") from exc

    rows: list[ImportRow] = []
    issues: list[RowIssue] = []
    issue_count = 0
    duplicate_count = 0
    total_populated = 0
    counts: dict[str, int] = {}
    seen_paths: dict[tuple[str | None, ...], tuple[str, int]] = {}
    topic_outcomes: dict[tuple[str | None, ...], str] = {}

    def issue(sheet: str, row: int, column: str | None, code: str) -> None:
        nonlocal issue_count
        issue_count += 1
        if issue_count > MAX_ISSUES:
            raise WorkbookRejected("too_many_validation_errors")
        issues.append(RowIssue(sheet, row, column, code))

    try:
        for sheet_name, expected_level in sheet_levels.items():
            if sheet_name not in workbook.sheetnames:
                raise WorkbookRejected("sheet_not_found")
            sheet = workbook[sheet_name]
            if sheet.sheet_state != "visible":
                raise WorkbookRejected("hidden_sheet")
            max_row, max_column = sheet.max_row, sheet.max_column
            if max_row > MAX_SHEET_ROWS or max_column > 100:
                raise WorkbookRejected("sheet_dimensions_too_large")
            if sheet.merged_cells.ranges:
                raise WorkbookRejected("merged_cells")
            headers = tuple(sheet.cell(1, col).value for col in range(1, max_column + 1))
            while headers and headers[-1] is None:
                headers = headers[:-1]
            if headers == HEADERS:
                field_columns = (1, 2, 3, 4, 5, 6, 7, 8, 9)
            elif headers == OPTIONAL_OUTCOME_HEADERS:
                field_columns = (1, 2, 3, 4, 5, 6, 7, None, 8)
            else:
                raise WorkbookRejected("invalid_headers")
            populated = 0
            blank_start: int | None = None
            for row_number in range(2, max_row + 1):
                cells = [sheet.cell(row_number, col) for col in range(1, max_column + 1)]
                has_data = any(cell.value is not None and str(cell.value).strip() for cell in cells)
                if not has_data:
                    if populated and blank_start is None:
                        blank_start = row_number
                    continue
                if blank_start is not None:
                    issue(sheet_name, blank_start, None, "internal_blank_row")
                    blank_start = None
                populated += 1
                total_populated += 1
                if total_populated > MAX_ROWS:
                    raise WorkbookRejected("too_many_rows")
                values: list[str | None] = []
                row_error = False
                for index, column in enumerate(field_columns):
                    if column is None:
                        values.append(None)
                        continue
                    cell = sheet.cell(row_number, column)
                    value = cell.value
                    if value is None or value == "":
                        values.append(None)
                    elif cell.data_type == "f":
                        issue(sheet_name, row_number, HEADERS[index], "formula_not_allowed")
                        values.append(None)
                        row_error = True
                    elif cell.data_type == "e":
                        issue(sheet_name, row_number, HEADERS[index], "error_cell")
                        values.append(None)
                        row_error = True
                    elif not isinstance(value, str):
                        issue(sheet_name, row_number, HEADERS[index], "text_required")
                        values.append(None)
                        row_error = True
                    else:
                        text = normalize_text(value)
                        if len(text) > MAX_FIELD_LENGTHS[index] or len(text.casefold()) > MAX_FIELD_LENGTHS[index]:
                            issue(sheet_name, row_number, HEADERS[index], "value_too_long")
                            row_error = True
                        values.append(text or None)
                for index in REQUIRED:
                    if values[index] is None:
                        issue(sheet_name, row_number, HEADERS[index], "required")
                        row_error = True
                if values[0] and match_key(values[0]) != match_key(expected_code):
                    issue(sheet_name, row_number, HEADERS[0], "curriculum_code_mismatch")
                    row_error = True
                if (expected_level != ALL_LEVELS and values[1] and
                        match_key(values[1]) != match_key(expected_level)):
                    issue(sheet_name, row_number, HEADERS[1], "sheet_level_mismatch")
                    row_error = True
                for cell in cells[len(headers):]:
                    if cell.value is not None and str(cell.value).strip():
                        issue(sheet_name, row_number, None, "extra_column")
                        row_error = True
                        break
                if row_error:
                    continue
                key = tuple(match_key(value) for value in values[:7])
                outcome = match_key(values[7])
                if outcome:
                    earlier = topic_outcomes.get(key[:6])
                    if earlier is not None and earlier != outcome:
                        issue(sheet_name, row_number, HEADERS[7], "outcome_conflict")
                        continue
                    topic_outcomes[key[:6]] = outcome
                original = seen_paths.get(key)
                if original is not None:
                    duplicate_count += 1
                else:
                    seen_paths[key] = (sheet_name, row_number)
                rows.append(ImportRow(sheet_name, row_number, tuple(values), key, original))
            if populated == 0:
                raise WorkbookRejected("empty_sheet")
            counts[sheet_name] = populated
    finally:
        workbook.close()
    return ParsedWorkbook(sha256(payload).hexdigest(), tuple(rows), tuple(issues),
                          issue_count, duplicate_count, counts)

