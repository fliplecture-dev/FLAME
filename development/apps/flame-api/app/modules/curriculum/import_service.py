"""Internal, admin-gated curriculum workbook import workflow.

No HTTP route calls this service until B3 authentication is available. Each
method owns a transaction; callers provide a session factory, not a live session.
"""

from collections import Counter
import base64
import binascii
from datetime import datetime, timedelta, timezone
from hashlib import sha256
import json
from io import BytesIO
from openpyxl import load_workbook
from uuid import UUID, uuid4

from sqlalchemy import and_, func, or_, select, update
from sqlalchemy.orm import Session

from app.modules.curriculum.drafts import replace_draft, published_paths, published_revision, _path_key
from app.modules.curriculum.import_models import (
    CurriculumImport, CurriculumImportCommit, CurriculumImportSource,
)
from app.modules.curriculum.models import (
    Category, Curriculum, CurriculumDraft, CurriculumRevision, Domain, Level, Subject, Subtopic, Topic,
)
from app.modules.curriculum.workbook_parser import (
    ALL_LEVELS, WorkbookRejected, _check_zip, match_key, parse_workbook,
)
from app.modules.identity.models import AdminRole


class ImportFailure(ValueError):
    def __init__(self, code: str, revision_mismatch: str | None = None):
        self.code = code
        self.revision_mismatch = revision_mismatch
        super().__init__(code)


def _require_admin(session: Session, owner: str) -> None:
    if not owner or session.get(AdminRole, owner) is None:
        raise ImportFailure("admin_required")


def _owned_import(session: Session, import_id: UUID, owner: str, *, lock: bool = False) -> CurriculumImport:
    query = select(CurriculumImport).where(
        CurriculumImport.id == import_id,
        CurriculumImport.owner_subject == owner,
    )
    if lock:
        query = query.with_for_update()
    batch = session.scalar(query)
    if batch is None:
        raise ImportFailure("import_not_found")
    return batch


def _existing_paths(session: Session, curriculum_id: UUID) -> tuple[set[tuple], dict[tuple, str | None]]:
    """Load all path ancestors in six scoped queries, including null categories."""
    levels = {node.id: node for node in session.scalars(
        select(Level).where(Level.curriculum_id == curriculum_id))}
    if not levels:
        return set(), {}
    subjects = {node.id: node for node in session.scalars(
        select(Subject).where(Subject.level_id.in_(levels)))}
    if not subjects:
        return set(), {}
    domains = {node.id: node for node in session.scalars(
        select(Domain).where(Domain.subject_id.in_(subjects)))}
    if not domains:
        return set(), {}
    categories = {node.id: node for node in session.scalars(
        select(Category).where(Category.domain_id.in_(domains)))}
    topics = {node.id: node for node in session.scalars(
        select(Topic).where(Topic.domain_id.in_(domains)))}
    if not topics:
        return set(), {}
    subtopics = session.scalars(select(Subtopic).where(Subtopic.topic_id.in_(topics))).all()
    paths: set[tuple] = set()
    outcomes: dict[tuple, str | None] = {}
    for topic in topics.values():
        domain = domains[topic.domain_id]
        subject = subjects[domain.subject_id]
        level = levels[subject.level_id]
        category_name = categories[topic.category_id].normalized_name if topic.category_id else None
        topic_path = (level.normalized_name, subject.normalized_name,
                      domain.normalized_name, category_name, topic.normalized_name)
        outcomes[topic_path] = match_key(topic.learning_outcome)
        paths.add(topic_path + (None,))
    for subtopic in subtopics:
        topic = topics[subtopic.topic_id]
        domain = domains[topic.domain_id]
        subject = subjects[domain.subject_id]
        level = levels[subject.level_id]
        category_name = categories[topic.category_id].normalized_name if topic.category_id else None
        paths.add((level.normalized_name, subject.normalized_name, domain.normalized_name,
                   category_name, topic.normalized_name, subtopic.normalized_name))
    return paths, outcomes


def _dependency_warnings(session, current_paths, incoming):
    from app.modules.learning.models import LearningMembership
    from app.modules.resources.models import ResourceTopic
    from app.modules.activity_kits.models import KitTopic
    removed = [row for row in current_paths if _path_key(row["values"]) not in incoming]
    if not removed:
        return []
    topics = {UUID(row["ids"]["topic"]) for row in removed if row.get("ids", {}).get("topic")
              and not any(path[:5] == _path_key(row["values"])[:5] for path in incoming)}
    subtopics = {UUID(row["ids"]["subtopic"]) for row in removed if row.get("ids", {}).get("subtopic")}
    resources = session.scalar(select(func.count()).select_from(ResourceTopic).where(
        or_(ResourceTopic.topic_id.in_(topics), ResourceTopic.subtopic_id.in_(subtopics)))) or 0
    memberships = session.scalar(select(func.count()).select_from(LearningMembership).where(
        or_(LearningMembership.topic_id.in_(topics), LearningMembership.subtopic_id.in_(subtopics)))) or 0
    kits = session.scalar(select(func.count()).select_from(KitTopic).where(KitTopic.topic_id.in_(topics))) or 0
    warnings = ["Missing paths are removed from the replacement draft. Publication archives old node IDs and preserves historical learning records."]
    if resources or memberships or kits:
        warnings.append(f"Removed paths have {resources} resource links, {memberships} My Learning links and {kits} activity-kit links. These paths become unavailable in the new member revision after publication; linked records are retained.")
    return warnings


class CurriculumImportService:
    def __init__(self, session_factory):
        self.session_factory = session_factory

    def upload(self, *, owner: str, curriculum_id: UUID, filename: str, payload: bytes) -> dict:
        if not filename or len(filename) > 255 or "/" in filename or "\\" in filename:
            raise ImportFailure("invalid_filename")
        if not filename.casefold().endswith(".xlsx"):
            raise WorkbookRejected("unsupported_file_type")
        _check_zip(payload)
        try:
            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            sheets = list(workbook.sheetnames)
            metadata = {}
            if "Metadata" in workbook.sheetnames:
                metadata = {
                    str(row[0]): str(row[1])
                    for row in workbook["Metadata"].iter_rows(min_row=1, max_col=2, values_only=True)
                    if row[0] is not None and row[1] is not None
                }
            workbook.close()
        except Exception as exc:
            raise WorkbookRejected("invalid_workbook") from exc
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            curriculum = session.get(Curriculum, curriculum_id)
            if curriculum is None or curriculum.state == "archived":
                raise ImportFailure("curriculum_not_found")
            if metadata:
                if metadata.get("schema_version") != "flame-curriculum-1":
                    raise ImportFailure("unsupported_workbook_schema")
                if (metadata.get("curriculum_id") != str(curriculum.id)
                        or match_key(metadata.get("curriculum_code")) != match_key(curriculum.code)):
                    raise ImportFailure("workbook_identity_mismatch")
            # A new upload supersedes this administrator's older pending workbook.
            for prior in session.scalars(select(CurriculumImport).where(
                    CurriculumImport.curriculum_id == curriculum_id,
                    CurriculumImport.owner_subject == owner,
                    CurriculumImport.status != "committed").with_for_update()):
                prior.preview_result = {**(prior.preview_result or {}), "cancelled": True,
                    "events": [*(prior.preview_result or {}).get("events", []), {"status": "superseded", "at": datetime.now(timezone.utc).isoformat()}]}
            batch = CurriculumImport(
                curriculum_id=curriculum_id, owner_subject=owner, filename=filename,
                file_hash=sha256(payload).hexdigest(), file_bytes=payload,
                status="uploaded", issue_count=0,
                preview_result={"events": [{"status": "uploaded", "at": datetime.now(timezone.utc).isoformat()}]},
                payload_expires_at=datetime.now(timezone.utc) + timedelta(days=30),
            )
            session.add(batch)
            session.flush()
            result = {"import_id": str(batch.id), "curriculum_id": str(curriculum_id),
                      "file_hash": batch.file_hash, "status": batch.status, "sheets": sheets,
                      "metadata": metadata, "round_trip": bool(metadata)}
        return result

    def preview(self, *, owner: str, import_id: UUID, sheet_levels: dict[str, str]) -> dict:
        with self.session_factory() as session, session.begin():
            return self._preview(session, owner=owner, import_id=import_id, sheet_levels=sheet_levels)

    def _preview(self, session, *, owner, import_id, sheet_levels, event_status="validated"):
        _require_admin(session, owner)
        batch = _owned_import(session, import_id, owner, lock=True)
        if (batch.preview_result or {}).get("cancelled"):
            raise ImportFailure("import_cancelled")
        if batch.status == "committed":
            raise ImportFailure("already_committed")
        if batch.file_bytes is None:
            raise ImportFailure("payload_expired")
        curriculum = session.scalar(select(Curriculum).where(
            Curriculum.id == batch.curriculum_id).with_for_update(read=True))
        if curriculum is None or curriculum.state == "archived":
            raise ImportFailure("curriculum_not_found")
        parsed = parse_workbook(batch.file_bytes, filename=batch.filename,
                                sheet_levels=sheet_levels, expected_code=curriculum.code)
        if parsed.file_hash != batch.file_hash:
            raise ImportFailure("file_hash_mismatch")
        draft = session.get(CurriculumDraft, curriculum.id)
        current_paths = draft.paths if draft else published_paths(session, curriculum.id)
        existing = {_path_key(row["values"]) for row in current_paths}
        _, existing_outcomes = _existing_paths(session, batch.curriculum_id)
        seen: set[tuple] = set()
        incoming_outcomes = {row.path_key[1:6]: match_key(row.values[7])
                             for row in parsed.rows if row.values[7]}
        counts = Counter(created=0, reused=0, updated=0, removed=0,
                         invalid=len({(e.sheet, e.row) for e in parsed.issues}),
                         duplicates=parsed.duplicate_count)
        preview_rows = []
        display_rows = []
        for row in parsed.rows:
            path = row.path_key[1:]
            topic_path = path[:5]
            if row.duplicate_of:
                operation = "duplicate"
            elif path in existing:
                outcome = incoming_outcomes.get(topic_path)
                operation = "updated" if outcome and outcome != existing_outcomes.get(topic_path) else "reused"
                counts[operation] += 1
            elif path in seen:
                operation = "duplicate"
            else:
                operation = "created"
                counts[operation] += 1
            seen.add(path)
            preview_rows.append({"sheet": row.sheet, "row": row.row,
                                 "values": list(row.values), "operation": operation})
            if len(display_rows) < 100:
                display_rows.append({"sheet": row.sheet, "row_number": row.row,
                                     "normalized_path": list(path), "operation": operation,
                                     "errors": []})
        removed_paths = sorted((path for path in existing
            if path[-1] is not None and path not in seen), key=lambda path: tuple(value or "" for value in path))
        counts["removed"] = len(removed_paths)
        for path in removed_paths:
            if len(display_rows) >= 100:
                break
            display_rows.append({
                "sheet": "Current curriculum", "row_number": 0,
                "normalized_path": list(path), "operation": "removed",
                "errors": [],
            })
        error_rows = [{"sheet": error.sheet, "row_number": error.row,
                       "column": error.column, "code": error.code}
                      for error in parsed.issues]
        if error_rows:
            display_rows = [{"sheet": error["sheet"], "row_number": error["row_number"],
                             "normalized_path": None, "operation": "invalid",
                             "errors": [{"column": error["column"], "code": error["code"]}]}
                            for error in error_rows[:100]] + display_rows[:max(0, 100 - len(error_rows))]
        revision = uuid4()
        batch.sheet_levels = dict(sheet_levels)
        batch.preview_revision = revision
        batch.base_revision = curriculum.revision
        batch.preview_rows = preview_rows
        batch.issue_count = parsed.issue_count
        batch.status = "previewed"
        events = list((batch.preview_result or {}).get("events", []))
        events.append({"status": event_status, "at": datetime.now(timezone.utc).isoformat(),
                       "curriculum_revision": curriculum.revision, "draft_revision": draft.revision if draft else 0,
                       "preview_revision": str(revision)})
        batch.preview_result = {
            "events": events, "removed_rows": [row["values"] for row in current_paths if _path_key(row["values"]) in removed_paths],
            "import_id": str(batch.id), "curriculum_id": str(batch.curriculum_id),
            "preview_revision": str(revision), "base_revision": curriculum.revision,
            "base_draft_revision": draft.revision if draft else 0,
            "curriculum_version": curriculum.version,
            "published_revision": published_revision(session, curriculum),
            "draft_revision": draft.revision if draft else 0,
            "dependency_warnings": _dependency_warnings(session, current_paths, seen),
            "file_hash": batch.file_hash, "status": "previewed",
            "counts": dict(counts), "sheet_counts": parsed.sheet_counts,
            "rows": display_rows, "issue_count": parsed.issue_count,
            "errors": error_rows,
        }
        result = {key: value for key, value in batch.preview_result.items() if key not in {"errors", "removed_rows", "events"}}
        return result

    def record_failure(self, *, owner, import_id, operation, code):
        with self.session_factory() as session, session.begin():
            try:
                _require_admin(session, owner)
                batch = _owned_import(session, import_id, owner, lock=True)
            except ImportFailure:
                return
            batch.preview_result = {**(batch.preview_result or {}), "events": [*(batch.preview_result or {}).get("events", []),
                {"status": operation + "_failed", "code": code, "at": datetime.now(timezone.utc).isoformat()}]}

    def cancel(self, *, owner: str, import_id: UUID) -> dict:
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner, lock=True)
            if batch.status == "committed":
                raise ImportFailure("already_committed")
            batch.preview_result = {**(batch.preview_result or {}), "cancelled": True,
                "events": [*(batch.preview_result or {}).get("events", []), {"status": "cancelled", "at": datetime.now(timezone.utc).isoformat()}]}
            return {"import_id": str(batch.id), "status": "cancelled"}

    @staticmethod
    def _check_preview_base(session, batch, curriculum):
        draft = session.scalar(select(CurriculumDraft).where(
            CurriculumDraft.curriculum_id == curriculum.id).with_for_update())
        saved = batch.preview_result or {}
        if saved.get("cancelled"):
            raise ImportFailure("import_cancelled")
        # Legacy previews lack draft/version tokens and must be regenerated.
        if curriculum.revision != batch.base_revision:
            raise ImportFailure("stale_preview", "import_base_revision")
        if saved.get("curriculum_version") != curriculum.version:
            raise ImportFailure("stale_preview", "curriculum_version")
        if saved.get("base_draft_revision") != (draft.revision if draft else 0):
            raise ImportFailure("stale_preview", "draft_revision")

    def commit(self, *, owner: str, import_id: UUID, preview_revision: UUID,
               idempotency_key: str, confirm: str) -> dict:
        if not idempotency_key or not idempotency_key.strip() or len(idempotency_key) > 255:
            raise ImportFailure("invalid_idempotency_key")
        request_hash = sha256(json.dumps([str(import_id), str(preview_revision), confirm],
                                         separators=(",", ":")).encode()).hexdigest()
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            # Serialize same-principal commits, including different keys, before replay checks.
            session.execute(select(AdminRole).where(AdminRole.auth_subject == owner).with_for_update()).scalar_one()
            prior = session.scalar(select(CurriculumImportCommit).where(
                CurriculumImportCommit.owner_subject == owner,
                CurriculumImportCommit.idempotency_key == idempotency_key,
            ))
            if prior is not None:
                if prior.request_hash != request_hash:
                    raise ImportFailure("idempotency_conflict")
                return dict(prior.result)
            batch = _owned_import(session, import_id, owner, lock=True)
            if batch.status != "previewed" or batch.preview_revision != preview_revision:
                raise ImportFailure("invalid_preview", "import_preview_revision")
            if batch.issue_count:
                raise ImportFailure("invalid_rows")
            curriculum = session.scalar(select(Curriculum).where(Curriculum.id == batch.curriculum_id).with_for_update())
            if curriculum is None or curriculum.state == "archived":
                raise ImportFailure("curriculum_not_found")
            self._check_preview_base(session, batch, curriculum)
            if any(not row["values"][6] for row in batch.preview_rows or []):
                raise ImportFailure("subtopic_required")
            if confirm.strip() != "REPLACE " + curriculum.code:
                raise ImportFailure("replacement_confirmation_mismatch")
            replacement_rows = [{
                "level": record["values"][1], "subject": record["values"][2],
                "domain": record["values"][3], "category": record["values"][4] or "",
                "topic": record["values"][5], "subtopic": record["values"][6],
            } for record in batch.preview_rows or []
               if record["operation"] != "duplicate"]
            draft = replace_draft(
                session, curriculum, replacement_rows, owner,
                source_import_id=batch.id)
            result = {"import_id": str(batch.id), "curriculum_id": str(curriculum.id),
                      "preview_revision": str(preview_revision),
                      "published_revision": published_revision(session, curriculum),
                      "draft_revision": draft.revision,
                      "status": "replacement_draft_created",
                      "counts": dict(batch.preview_result["counts"])}
            batch.status = "committed"
            batch.committed_at = datetime.now(timezone.utc)
            batch.committed_result = result
            batch.preview_result = {**(batch.preview_result or {}), "events": [*(batch.preview_result or {}).get("events", []),
                {"status": "replacement_draft_created", "at": batch.committed_at.isoformat(), "draft_revision": draft.revision,
                 "curriculum_revision": curriculum.revision, "preview_revision": str(preview_revision)}]}
            session.add(CurriculumImportCommit(import_id=batch.id, owner_subject=owner,
                                               idempotency_key=idempotency_key,
                                               request_hash=request_hash, result=result))
        return result

    def _apply_rows(self, session: Session, batch: CurriculumImport,
                    curriculum: Curriculum) -> tuple[int, int, int, int]:
        caches: dict[type, dict[tuple, object]] = {cls: {} for cls in (Level, Subject, Domain, Category, Topic, Subtopic)}
        next_orders: dict[tuple, int] = {}
        seen_sources: set[tuple] = set()
        created = reused = updated = source_added = 0

        def get_or_create(cls, parent_field: str, parent_id: UUID, name: str, **kwargs):
            nonlocal created, reused
            normalized = match_key(name)
            category_id = kwargs.get("category_id") if cls is Topic else None
            key = (parent_id, category_id, normalized) if cls is Topic else (parent_id, normalized)
            cache = caches[cls]
            if key in cache:
                return cache[key]
            query = select(cls).where(getattr(cls, parent_field) == parent_id,
                                      cls.normalized_name == normalized)
            if cls is Topic:
                query = query.where(cls.category_id.is_(None) if category_id is None else cls.category_id == category_id)
            node = session.scalar(query)
            if node is None:
                order_key = (cls.__name__, parent_id, kwargs.get("category_id"))
                if order_key not in next_orders:
                    order_query = select(func.max(cls.display_order)).where(getattr(cls, parent_field) == parent_id)
                    if cls is Topic:
                        category_id = kwargs.get("category_id")
                        order_query = order_query.where(cls.category_id.is_(None) if category_id is None else cls.category_id == category_id)
                    maximum = session.scalar(order_query)
                    next_orders[order_key] = 0 if maximum is None else maximum + 1
                order = next_orders[order_key]
                next_orders[order_key] += 1
                data = dict(name=name, normalized_name=match_key(name), display_order=order,
                            **{parent_field: parent_id}, **kwargs)
                if cls is Level:
                    data.update(code=name, normalized_code=match_key(name))
                node = cls(**data)
                session.add(node)
                session.flush()
                created += 1
            else:
                if node.archived:
                    raise ImportFailure("archived_path_conflict")
                reused += 1
            cache[key] = node
            return node

        for record in batch.preview_rows or []:
            values = record["values"]
            _, level_name, subject_name, domain_name, category_name, topic_name, subtopic_name, outcome, source = values
            level = get_or_create(Level, "curriculum_id", curriculum.id, level_name)
            subject = get_or_create(Subject, "level_id", level.id, subject_name)
            domain = get_or_create(Domain, "subject_id", subject.id, domain_name)
            category = (get_or_create(Category, "domain_id", domain.id, category_name)
                        if category_name else None)
            topic = get_or_create(Topic, "domain_id", domain.id, topic_name,
                                  category_id=category.id if category else None)
            if outcome and topic.learning_outcome != outcome:
                topic.learning_outcome = outcome
                topic.version += 1
                updated += 1
            subtopic = (get_or_create(Subtopic, "topic_id", topic.id, subtopic_name)
                        if subtopic_name else None)
            if source:
                source_key = (topic.id, subtopic.id if subtopic else None, source)
                if source_key not in seen_sources:
                    existing = session.scalar(select(CurriculumImportSource).where(
                        CurriculumImportSource.topic_id == topic.id,
                        CurriculumImportSource.subtopic_id.is_(None) if subtopic is None else CurriculumImportSource.subtopic_id == subtopic.id,
                        CurriculumImportSource.source_text == source,
                    ))
                    if existing is None:
                        session.add(CurriculumImportSource(curriculum_id=curriculum.id,
                                                           topic_id=topic.id,
                                                           subtopic_id=subtopic.id if subtopic else None,
                                                           source_text=source,
                                                           first_import_id=batch.id))
                        source_added += 1
                    seen_sources.add(source_key)
        return created, reused, updated, source_added

    def detail(self, *, owner: str, import_id: UUID) -> dict:
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner)
            sheets = []
            metadata = {}
            if batch.file_bytes is not None:
                workbook = load_workbook(BytesIO(batch.file_bytes), read_only=True, data_only=True)
                try:
                    sheets = workbook.sheetnames
                    if "Metadata" in sheets:
                        metadata = {str(row[0]): str(row[1]) for row in workbook["Metadata"].iter_rows(max_col=2, values_only=True)
                                    if row[0] is not None and row[1] is not None}
                finally:
                    workbook.close()
            return {"import_id": str(batch.id), "curriculum_id": str(batch.curriculum_id),
                    "status": batch.status, "file_hash": batch.file_hash,
                    "filename": batch.filename, "sheets": sheets, "sheet_levels": batch.sheet_levels,
                    "round_trip": bool(metadata), "metadata": metadata,
                    "preview": ({key: value for key, value in batch.preview_result.items()
                                 if key not in {"errors", "removed_rows", "events"}} if batch.status == "previewed" and batch.preview_revision else None),
                    "committed_result": batch.committed_result, "events": (batch.preview_result or {}).get("events", []),
                    "payload_expires_at": batch.payload_expires_at.isoformat()}

    def issue_page(self, *, owner: str, import_id: UUID, offset: int = 0,
                   limit: int = 100) -> dict:
        if offset < 0 or not 1 <= limit <= 100:
            raise ImportFailure("invalid_page")
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner)
            errors = (batch.preview_result or {}).get("errors", [])
            page = errors[offset:offset + limit]
            return {"items": page, "total": batch.issue_count,
                    "next_offset": offset + len(page) if offset + len(page) < len(errors) else None}

    def row_page(self, *, owner: str, import_id: UUID, offset: int = 0,
                 limit: int = 100) -> dict:
        if offset < 0 or not 1 <= limit <= 100:
            raise ImportFailure("invalid_page")
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner)
            all_rows = batch.preview_rows or []
            page = all_rows[offset:offset + limit]
            return {"items": [{"sheet": row["sheet"], "row_number": row["row"],
                               "normalized_path": [match_key(v) for v in row["values"][1:7]],
                               "operation": row["operation"]} for row in page],
                    "total": len(all_rows),
                    "next_offset": offset + len(page) if offset + len(page) < len(all_rows) else None}

    def source_page(self, *, owner: str, import_id: UUID, offset: int = 0,
                    limit: int = 100) -> dict:
        """Read uploaded workbook rows without applying them to a curriculum."""
        if offset < 0 or not 1 <= limit <= 100:
            raise ImportFailure("invalid_page")
        with self.session_factory() as session:
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner)
            if batch.file_bytes is None:
                raise ImportFailure("payload_expired")
            curriculum = session.get(Curriculum, batch.curriculum_id)
            if curriculum is None:
                raise ImportFailure("curriculum_not_found")
            workbook = load_workbook(BytesIO(batch.file_bytes), read_only=True, data_only=True)
            rows = []
            workbook_codes = set()
            try:
                for sheet in workbook.worksheets:
                    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
                        if not any(value is not None and str(value).strip() for value in cells):
                            continue
                        values = ["" if value is None else str(value) for value in cells]
                        if values[0].strip():
                            workbook_codes.add(values[0].strip())
                        rows.append({"sheet": sheet.title, "row_number": row_number, "values": values})
            finally:
                workbook.close()
            page = rows[offset:offset + limit]
            return {"items": page, "total": len(rows),
                    "next_offset": offset + len(page) if offset + len(page) < len(rows) else None,
                    "workbook_codes": sorted(workbook_codes)[:10], "expected_code": curriculum.code}

    def review_page(self, *, owner: str, import_id: UUID, offset: int = 0,
                    limit: int = 100, operation: str = "", q: str = "") -> dict:
        if offset < 0 or not 1 <= limit <= 100:
            raise ImportFailure("invalid_page")
        with self.session_factory() as session:
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner)
            if batch.status != "previewed" or batch.file_bytes is None:
                raise ImportFailure("invalid_preview", "import_preview_revision")
            workbook = load_workbook(BytesIO(batch.file_bytes), read_only=True, data_only=True)
            operations = {(row["sheet"], row["row"]): row["operation"]
                         for row in batch.preview_rows or []}
            errors = {}
            for item in (batch.preview_result or {}).get("errors", []):
                errors.setdefault((item["sheet"], item["row_number"]), []).append(
                    {"column": item["column"], "code": item["code"]})
            all_rows = []
            try:
                for sheet_name in (batch.sheet_levels or {}):
                    sheet = workbook[sheet_name]
                    for row_number, values in enumerate(
                            sheet.iter_rows(min_row=2, max_col=9, values_only=True), start=2):
                        key = (sheet_name, row_number)
                        if not any(value is not None and str(value).strip() for value in values) and key not in errors:
                            continue
                        all_rows.append({"sheet": sheet_name, "row_number": row_number,
                                         "values": ["" if value is None else str(value)
                                                    for value in values[1:7]],
                                         "operation": operations.get(key, "invalid"),
                                         "errors": errors.get(key, [])})
            finally:
                workbook.close()
            uploaded_total = len(all_rows)
            if operation == "removed":
                all_rows = [{"sheet": "Saved draft", "row_number": index + 1,
                    "values": [values.get(field, "") for field in ("level", "subject", "domain", "category", "topic", "subtopic")],
                    "operation": "removed", "errors": []} for index, values in enumerate((batch.preview_result or {}).get("removed_rows", []))]
            elif operation:
                all_rows = [row for row in all_rows if (row["errors"] if operation == "invalid" else
                    row["operation"] in {"created", "updated"} if operation == "changed" else row["operation"] == operation)]
            if q:
                all_rows = [row for row in all_rows if q.casefold() in " / ".join(row["values"]).casefold()]
            page = all_rows[offset:offset + limit]
            return {"items": page, "total": len(all_rows), "uploaded_total": uploaded_total,
                    "next_offset": offset + len(page) if offset + len(page) < len(all_rows) else None}

    def review_changes(self, *, owner: str, import_id: UUID, preview_revision: UUID,
                       changes: list[dict]) -> dict:
        if not 1 <= len(changes) <= 100:
            raise ImportFailure("invalid_review_changes")
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            batch = _owned_import(session, import_id, owner, lock=True)
            if batch.status != "previewed" or batch.preview_revision != preview_revision:
                raise ImportFailure("invalid_preview", "import_preview_revision")
            if batch.file_bytes is None:
                raise ImportFailure("payload_expired")
            curriculum = session.scalar(select(Curriculum).where(
                Curriculum.id == batch.curriculum_id).with_for_update())
            if curriculum is None or curriculum.state == "archived":
                raise ImportFailure("curriculum_not_found")
            self._check_preview_base(session, batch, curriculum)
            workbook = load_workbook(BytesIO(batch.file_bytes), keep_links=False)
            deletions = {}
            additions = {}
            row_remap = {}
            originals = {name: workbook[name].max_row for name in batch.sheet_levels or {}}
            try:
                for change in changes:
                    sheet_name = change["sheet"]
                    if sheet_name not in (batch.sheet_levels or {}):
                        raise ImportFailure("sheet_not_mapped")
                    sheet = workbook[sheet_name]
                    action = change["action"]
                    row_number = change.get("row_number")
                    if action == "add":
                        row_number = change.get("insert_before") or originals[sheet_name] + 1
                        if not 2 <= row_number <= originals[sheet_name] + 1:
                            raise ImportFailure("row_not_found")
                    elif row_number is None or row_number < 2 or row_number > sheet.max_row:
                        raise ImportFailure("row_not_found")
                    if action == "delete":
                        deletions.setdefault(sheet_name, set()).add(row_number)
                        continue
                    values = change.get("values")
                    if values is None or len(values) != 6:
                        raise ImportFailure("invalid_review_values")
                    if (batch.sheet_levels[sheet_name] != ALL_LEVELS and
                            match_key(values[0]) != match_key(batch.sheet_levels[sheet_name])):
                        raise ImportFailure("level_must_match_sheet_mapping")
                    if action == "add":
                        additions.setdefault(sheet_name, {}).setdefault(row_number, []).append([curriculum.code, *values])
                    else:
                        sheet.cell(row_number, 1).value = curriculum.code
                        for index, value in enumerate(values, start=2):
                            sheet.cell(row_number, index).value = value.strip() or None
                for sheet_name, maximum in originals.items():
                    sheet = workbook[sheet_name]
                    deleted = deletions.get(sheet_name, set())
                    inserted = additions.get(sheet_name, {})
                    shift = 0
                    for row_number in range(2, maximum + 2):
                        shift += len(inserted.get(row_number, []))
                        row_remap[sheet_name + ":" + str(row_number)] = None if row_number in deleted else row_number + shift
                        if row_number in deleted:
                            shift -= 1
                    for row_number in sorted(set(deleted) | set(inserted), reverse=True):
                        if row_number in deleted:
                            sheet.delete_rows(row_number)
                        for values in reversed(inserted.get(row_number, [])):
                            sheet.insert_rows(row_number)
                            for index, value in enumerate(values, start=1):
                                sheet.cell(row_number, index).value = value.strip() or None
                output = BytesIO()
                workbook.save(output)
                payload = output.getvalue()
            finally:
                workbook.close()
            _check_zip(payload)
            parse_workbook(payload, filename=batch.filename,
                           sheet_levels=batch.sheet_levels, expected_code=curriculum.code)
            batch.file_bytes = payload
            batch.file_hash = sha256(payload).hexdigest()
            result = self._preview(session, owner=owner, import_id=import_id,
                sheet_levels=dict(batch.sheet_levels), event_status="corrections_validated")
            result["row_remap"] = row_remap
            return result

    def history(self, *, owner: str, curriculum_id: UUID, limit: int = 25,
                cursor: str | None = None) -> dict:
        if not 1 <= limit <= 100:
            raise ImportFailure("invalid_limit")
        boundary = None
        if cursor:
            if len(cursor) > 512:
                raise ImportFailure("invalid_cursor")
            try:
                raw = base64.urlsafe_b64decode(cursor + "=" * (-len(cursor) % 4))
                values = json.loads(raw)
                if not isinstance(values, list) or len(values) != 2:
                    raise ValueError
                boundary = (datetime.fromisoformat(values[0]), UUID(values[1]))
            except (ValueError, TypeError, OverflowError, binascii.Error) as exc:
                raise ImportFailure("invalid_cursor") from exc
        with self.session_factory() as session, session.begin():
            _require_admin(session, owner)
            query = select(CurriculumImport).where(
                CurriculumImport.owner_subject == owner,
                CurriculumImport.curriculum_id == curriculum_id,
            ).order_by(CurriculumImport.created_at.desc(), CurriculumImport.id.desc()).limit(limit + 1)
            if boundary is not None:
                when, identifier = boundary
                query = query.where(or_(CurriculumImport.created_at < when,
                                        and_(CurriculumImport.created_at == when,
                                             CurriculumImport.id < identifier)))
            batches = session.scalars(query).all()
            page = batches[:limit]
            next_cursor = None
            if len(batches) > limit:
                last = page[-1]
                raw = json.dumps([last.created_at.isoformat(), str(last.id)],
                                 separators=(",", ":")).encode()
                next_cursor = base64.urlsafe_b64encode(raw).decode().rstrip("=")
            return {"items": [{"import_id": str(b.id), "status": "cancelled" if (b.preview_result or {}).get("cancelled") else b.status,
                               "filename": b.filename, "file_hash": b.file_hash, "created_at": b.created_at.isoformat(),
                               "events": (b.preview_result or {}).get("events", []),
                               "draft_revision": (b.committed_result or b.preview_result or {}).get("draft_revision"),
                               "published_revision": (b.committed_result or b.preview_result or {}).get("published_revision"),
                               "counts": (b.committed_result or b.preview_result or {}).get("counts")}
                              for b in page], "next_cursor": next_cursor}

    def purge_expired_payloads(self, *, now: datetime | None = None) -> int:
        cutoff = now or datetime.now(timezone.utc)
        with self.session_factory() as session, session.begin():
            result = session.execute(update(CurriculumImport).where(
                CurriculumImport.payload_expires_at < cutoff,
                CurriculumImport.file_bytes.is_not(None),
            ).values(file_bytes=None))
            return result.rowcount






