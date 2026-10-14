"""Admin-only manual Sheets sync, with durable catalogue commit and retryable write-back."""
from datetime import datetime, timezone
from uuid import UUID
from io import BytesIO
from fastapi import APIRouter, Request
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import text,select
from app.modules.bulk_transfer.router import Admin, Confirm, principal_subject, locked_job, response, commit
from app.modules.bulk_transfer.models import BulkImport,GoogleSheetConnection
from app.modules.bulk_transfer.sheet_connector import connection,actions,sheet_snapshot,export_writeback
from app.modules.bulk_transfer.workbooks import workbook
from app.modules.bulk_transfer.template_guidance import resource_guidance, NOTES
from app.modules.bulk_transfer.service import digest
from app.modules.bulk_transfer.errors import fail, BulkFailure
from app.modules.bulk_transfer.google_sheets import SheetsClient, credentials, spreadsheet_id, collect, sheet_plan, recheck, writeback

router = APIRouter(prefix="/api/v1/admin/bulk/google-sheets", tags=["admin-google-sheets"])


class Preview(BaseModel):
    spreadsheet: str = Field(min_length=20, max_length=500)
    connector: bool = True


@router.get('/connections')
def connections(request:Request,principal:Admin):
    with request.app.state.database.session() as db:
        owner=principal_subject(db,principal)
        items=db.scalars(select(GoogleSheetConnection).where(GoogleSheetConnection.owner_subject==owner).order_by(GoogleSheetConnection.created_at.desc())).all()
        latest=db.scalar(select(BulkImport).where(BulkImport.owner_subject==owner,BulkImport.payload.has_key('sheets'),BulkImport.status.in_(['review','completed'])).order_by(BulkImport.created_at.desc()))
        pending=latest if latest and (latest.status=='review' or (latest.result or {}).get('writeback')=='pending') else None
        return {'items':[{'id':str(x.id),'spreadsheet_id':x.spreadsheet_id,'last_synced_at':x.last_synced_at} for x in items],'pending_job_id':str(pending.id) if pending else None}


@router.get("/template")
def template(request: Request, principal: Admin):
    book = load_workbook(BytesIO(workbook("resources")))
    del book["Mappings"]; del book["Paths"]
    book["Records"].title = "Image"
    for name in ("Video", "Interactive"): book.copy_worksheet(book["Image"]).title = name
    for name in ("Image", "Video", "Interactive"):
        book[name].cell(1,book[name].max_column+1,'sync_action')
        book[name].auto_filter.ref = book[name].dimensions
        book[name].sheet_view.showGridLines = True
        resource_guidance(book[name], name.lower())
    book["Guide"].append(["Google Sheets: keep Image, Video and Interactive tabs. New identifiers may be blank. Sync reviews metadata and identifiers; mappings/assets stay in FLAME. Keep linked keys/IDs and set sync_action=delete to request deletion. Confirm separately; FLAME then writes deleted. Deleted rows cannot be re-imported."])
    book["Guide"].append(["Header notes and data validation cover rows 2–2001. Pasting data or converting to Google Sheets may bypass/change validation; FLAME still validates every import. Examples below are instructions, not resource records."])
    book["Guide"].append(["Column", "What to enter", "Example"])
    for field,(instruction,example) in NOTES.items(): book["Guide"].append([field,instruction,example])
    book["Guide"].column_dimensions["B"].width = 80
    book["Guide"].column_dimensions["C"].width = 55
    out = BytesIO(); book.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="flame-google-sheets-template.xlsx"', "Cache-Control": "private, no-store"})


@router.get("/configuration")
def configuration(request: Request, principal: Admin):
    if not (request.app.state.settings.google_sheets_credentials_json or request.app.state.settings.google_sheets_credentials_file):
        return {"configured": False, "service_account_email": None}
    auth = credentials(request.app.state.settings)
    return {"configured": True, "service_account_email": auth.service_account_email}


@router.post("/preview")
def preview(request: Request, principal: Admin, body: Preview):
    config = request.app.state.settings
    client = SheetsClient(config, spreadsheet_id(body.spreadsheet))
    with request.app.state.database.session() as db, db.begin():
        owner = principal_subject(db, principal)
        # Serialize FLAME previews for a spreadsheet while reserving row metadata.
        db.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:key, 0))"), {"key": "sheets:" + client.ident})
        conn=connection(db,owner,client.ident,True) if body.connector else None
        from app.modules.bulk_transfer.sheet_lifecycle import bind_client
        bind_client(client,db,conn=conn)
        records, refs, issues, markers = collect(*client.read(), reserve=True, namespace=client.ident)
        if any(x["severity"] == "error" for x in issues): return {"job": None, "sheet_issues": issues}
        if not records and not body.connector: return {"job": None, "sheet_issues": issues + [{"severity":"error","tab":"","row":0,"column":"","message":"No resource rows found."}]}
        assigned = client.mark(markers)
        for ref in refs:
            if ref["metadata_id"] is None: ref["metadata_id"] = assigned[ref["marker"]]
        source = {"spreadsheet_id": client.ident, "rows": refs, "issues": issues}
        if conn:source['connection_id']=str(conn.id)
        recheck(client, source)
        payload = {"records": records, "mappings": [], "uploads": {}, "asset_names": [], "sheets": source}
        if conn:
            for record,ref in zip(records,refs):
                bound=conn.bindings.get(ref['marker'])
                if not bound:continue
                if record.get('existing_id') and str(record['existing_id'])!=bound['resource_id']:
                    ref['errors'].append('This tracked row belongs to another resource. Use a new row for a new resource; preserve linked identifiers.')
                record.setdefault('existing_id',bound['resource_id'])
                if not ref['source'].get('record_key'):
                    record['record_key']=bound['record_key'];ref['record_key']=bound['record_key']
        review = sheet_plan(db, config, payload)
        if body.connector:
            if db.scalar(select(BulkImport.id).where(BulkImport.owner_subject==owner,BulkImport.status=='completed',BulkImport.result['writeback'].astext=='pending',BulkImport.payload['sheets']['spreadsheet_id'].astext==client.ident)):
                fail(409,'This connection has a pending sheet update. Retry that saved sync before starting another.')
            for ref,entry in zip(refs,review['entries']):
                ref['matched_id']=entry.get('existing_id');ref['operation']=entry['operation']
            names,grids,tracked=client.read()
            for tab in ('Image','Video','Interactive'):
                if 'sync_action' not in grids[tab][0]:issues.append({'severity':'warning','tab':tab,'row':1,'column':'sync_action','message':'Confirmed sync will add the sync_action deletion column and its dropdown automatically. Existing columns are preserved.'})
            source['tracking_cleanup']=getattr(client,'tracking_cleanup',[])
            if source['tracking_cleanup']:issues.append({'severity':'warning','tab':'','row':0,'column':'','message':f"{len(source['tracking_cleanup'])} duplicate tracking markers can be repaired after sync confirmation; matched row identities are preserved."})
            source['connector']=actions(db,conn,names,grids,refs)
            source['sheet_snapshot']=sheet_snapshot(names,grids,tracked)
            review=sheet_plan(db,config,payload)
        job = BulkImport(kind="resources", owner_subject=owner, filename="Google Sheets · Image / Video / Interactive", payload=payload, preview=review, token=digest(review))
        db.add(job); db.flush()
        return {"job": response(job), "sheet_issues": issues}


@router.post("/{job_id}/writeback")
def retry_writeback(request: Request, principal: Admin, job_id: UUID):
    with request.app.state.database.session() as db, db.begin():
        job = locked_job(db, job_id, principal_subject(db, principal))
        if "sheets" not in job.payload or job.status != "completed": fail(409, "Complete this Google Sheets import before writing identifiers.")
        if job.result.get("writeback") == "completed": return response(job)
        try:
            client=SheetsClient(request.app.state.settings, job.payload['sheets']['spreadsheet_id'])
            from app.modules.bulk_transfer.sheet_lifecycle import bind_client
            bind_client(client,db,job.payload['sheets'])
            if 'connector' in job.payload['sheets']:
                export_writeback(client,db,job.payload['sheets'],job.result['ids'],job.result['sheet_records'])
            else:writeback(client, job.payload["sheets"], job.result["ids"])
            job.result = {**job.result, "writeback": "completed", "writeback_error": None, "writeback_at": datetime.now(timezone.utc).isoformat()}
        except BulkFailure as error:
            job.result = {**job.result, "writeback": "pending", "writeback_error": error.detail}
        return response(job)


@router.post("/{job_id}/commit")
def confirm(request: Request, principal: Admin, job_id: UUID, body: Confirm):
    with request.app.state.database.session() as db, db.begin():
        job = locked_job(db, job_id, principal_subject(db, principal))
        if "sheets" not in job.payload: fail(404, "Google Sheets sync not found.")
    commit(request, principal, job_id, body)
    # Database transaction completes first. A Google failure leaves durable pending state.
    return retry_writeback(request, principal, job_id)
