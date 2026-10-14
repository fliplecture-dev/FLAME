"""Admin-only private staging. Only small tickets/results cross hosted Functions."""
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Annotated, Literal
from uuid import UUID, uuid4
import re

import httpx
from cryptography.fernet import Fernet, InvalidToken
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from sqlalchemy.exc import IntegrityError
from starlette.datastructures import Headers

from app.modules.identity.dependencies import require_admin
from app.modules.identity.models import StudentProfile
from app.modules.resources.validation import MediaRejected, declared_format

router = APIRouter(prefix='/api/v1/admin/transfers', tags=['admin-private-transfers'])
Admin = Annotated[dict, Depends(require_admin)]
TICKET_TTL = 2 * 60 * 60


def fail(status, code):
    raise HTTPException(status, code)


class Admission(BaseModel):
    model_config = ConfigDict(extra='forbid')
    purpose: Literal['resource', 'curriculum', 'bulk']
    filename: str = Field(min_length=1, max_length=255)
    mime_type: str = Field(max_length=100)
    byte_size: int = Field(gt=0, le=120 * 1024 * 1024)
    sha256: str = Field(pattern=r'^[0-9a-f]{64}$')
    curriculum_id: UUID | None = None
    kind: Literal['resources', 'teachers', 'students', 'activity-kits'] | None = None

    @field_validator('filename')
    @classmethod
    def safe_filename(cls, value):
        if '/' in value or '\\' in value or any(ord(c) < 32 for c in value):
            raise ValueError('invalid_filename')
        return value

    @model_validator(mode='after')
    def limits(self):
        if (self.purpose == 'curriculum') != bool(self.curriculum_id) or (self.purpose == 'bulk') != bool(self.kind):
            raise ValueError('invalid_upload_target')
        if self.purpose == 'resource':
            try:
                _, _, _, limit = declared_format(self.filename, self.mime_type)
            except MediaRejected as error:
                raise ValueError(str(error)) from None
        elif self.purpose == 'curriculum':
            if not self.filename.lower().endswith('.xlsx'):
                raise ValueError('unsupported_file_type')
            limit = 10 * 1024 * 1024
        else:
            if Path(self.filename).suffix.lower() not in {'.xlsx', '.zip'}:
                raise ValueError('unsupported_file_type')
            limit = 120 * 1024 * 1024
        if self.byte_size > limit:
            raise ValueError('file_too_large')
        return self


class Ticket(BaseModel):
    model_config = ConfigDict(extra='forbid')
    version: Literal[1]
    id: UUID
    owner: str
    admission: Admission


class Completion(BaseModel):
    model_config = ConfigDict(extra='forbid')
    ticket: str = Field(min_length=1, max_length=8192)


class Export(BaseModel):
    model_config = ConfigDict(extra='forbid')
    purpose: Literal['bulk', 'curriculum']
    kind: Literal['resources', 'teachers', 'students', 'activity-kits'] | None = None
    action: Literal['template', 'export'] = 'export'
    package: bool = False
    curriculum_id: UUID | None = None
    source: Literal['draft', 'published', 'revision'] = 'draft'
    revision_number: int | None = Field(default=None, ge=1)

    @model_validator(mode='after')
    def target(self):
        if (self.purpose == 'curriculum') != bool(self.curriculum_id) or (self.purpose == 'bulk') != bool(self.kind):
            raise ValueError('invalid_export_target')
        return self


def cloud(config):
    secret = config.supabase_secret_key or config.supabase_service_role_key
    if (config.resource_storage_backend != 'supabase' or not config.resource_storage_bucket
            or not config.supabase_url or not secret or not config.session_encryption_key):
        fail(503, 'direct_storage_unavailable')
    if not re.fullmatch(r'[A-Za-z0-9_.-]{1,100}', config.resource_storage_bucket):
        fail(503, 'direct_storage_unavailable')
    token = secret.get_secret_value()
    return config.supabase_url + '/storage/v1', config.resource_storage_bucket, {
        'apikey': token, 'Authorization': 'Bearer ' + token}


def cipher(config):
    try:
        return Fernet(config.session_encryption_key.get_secret_value().encode())
    except (AttributeError, ValueError):
        fail(503, 'direct_storage_unavailable')


def owner(request, principal):
    with request.app.state.database.session() as db:
        profile = db.get(StudentProfile, UUID(principal['id']))
        if profile is None or profile.status != 'active':
            fail(403, 'account_inactive')
        return profile.auth_subject


def stage_key(ticket):
    return 'staged-uploads/' + ticket.id.hex


@router.get('/configuration')
def configuration(request: Request, _: Admin):
    config = request.app.state.settings
    secret = config.supabase_secret_key or config.supabase_service_role_key
    enabled = bool(config.resource_storage_backend == 'supabase' and config.resource_storage_bucket
                   and config.supabase_url and secret and config.session_encryption_key)
    return {'direct_upload': enabled, 'storage_origin': config.supabase_url if enabled else None}


@router.post('/authorize')
def authorize(request: Request, principal: Admin, body: Admission):
    config = request.app.state.settings
    origin, bucket, headers = cloud(config)
    subject = owner(request, principal)
    if body.curriculum_id:
        from app.modules.curriculum.models import Curriculum
        with request.app.state.database.session() as db:
            curriculum = db.get(Curriculum, body.curriculum_id)
            if curriculum is None or curriculum.state == 'archived':
                fail(404, 'curriculum_not_found')
    ticket = Ticket(version=1, id=uuid4(), owner=subject, admission=body)
    path = '/object/upload/sign/' + bucket + '/' + stage_key(ticket)
    try:
        with httpx.Client(timeout=20, follow_redirects=False) as client:
            response = client.post(origin + path, headers=headers, json={})
        if response.status_code != 200:
            fail(503, 'direct_storage_unavailable')
        url = response.json().get('url')
        if not isinstance(url, str) or not url.startswith(path + '?token='):
            fail(503, 'direct_storage_unavailable')
    except (httpx.HTTPError, ValueError, AttributeError):
        fail(503, 'direct_storage_unavailable')
    return {'ticket': cipher(config).encrypt(ticket.model_dump_json().encode()).decode(),
            'upload_url': origin + url, 'storage_origin': config.supabase_url,
            'expires_in': TICKET_TTL}


def existing_result(request, ticket):
    from app.modules.resources.models import ResourceUpload
    from app.modules.bulk_transfer.models import BulkImport
    from app.modules.bulk_transfer.router import response
    from app.modules.curriculum.import_models import CurriculumImport
    body = ticket.admission
    model = {'resource': ResourceUpload, 'bulk': BulkImport, 'curriculum': CurriculumImport}[body.purpose]
    with request.app.state.database.session() as db:
        row = db.get(model, ticket.id)
        if row is None:
            return None
        subject = row.owner_auth_subject if body.purpose == 'resource' else row.owner_subject
        if subject != ticket.owner:
            fail(404, 'upload_not_found')
        if body.purpose == 'resource':
            return {'upload_id': str(row.id), 'mime_type': row.mime_type,
                    'byte_size': row.byte_size, 'media_type': row.media_type}
        if body.purpose == 'bulk':
            if row.kind != body.kind:
                fail(409, 'upload_target_conflict')
            return response(row)
        if row.curriculum_id != body.curriculum_id:
            fail(409, 'upload_target_conflict')
        # Workbook bytes and preview state are already durable in PostgreSQL.
        from io import BytesIO
        from openpyxl import load_workbook
        sheets = []
        if row.file_bytes:
            book = load_workbook(BytesIO(row.file_bytes), read_only=True, data_only=True)
            sheets = list(book.sheetnames)
            book.close()
        return {'import_id': str(row.id), 'curriculum_id': str(row.curriculum_id),
                'file_hash': row.file_hash, 'status': row.status, 'sheets': sheets}


def download_stage(config, ticket, destination):
    origin, bucket, headers = cloud(config)
    total, digest = 0, sha256()
    try:
        with httpx.Client(timeout=120, follow_redirects=False) as client:
            with client.stream('GET', origin + '/object/' + bucket + '/' + stage_key(ticket), headers=headers) as response:
                if response.status_code != 200:
                    fail(409, 'staged_upload_missing')
                with destination.open('wb') as out:
                    for chunk in response.iter_bytes(65536):
                        total += len(chunk)
                        if total > ticket.admission.byte_size:
                            fail(413, 'upload_size_mismatch')
                        digest.update(chunk)
                        out.write(chunk)
    except httpx.HTTPError:
        fail(503, 'direct_storage_unavailable')
    if total != ticket.admission.byte_size or digest.hexdigest() != ticket.admission.sha256:
        fail(422, 'upload_content_mismatch')


def remove_stage(config, ticket):
    origin, bucket, headers = cloud(config)
    try:
        with httpx.Client(timeout=15, follow_redirects=False) as client:
            client.request('DELETE', origin + '/object/' + bucket, headers=headers,
                           json={'prefixes': [stage_key(ticket)]})
    except httpx.HTTPError:
        pass  # Receipt replay remains valid; stale stages are operations cleanup work.


@router.post('/complete')
def complete(request: Request, principal: Admin, body: Completion):
    config = request.app.state.settings
    cloud(config)
    try:
        ticket = Ticket.model_validate_json(cipher(config).decrypt(body.ticket.encode(), ttl=TICKET_TTL))
    except (InvalidToken, ValueError):
        fail(410, 'upload_ticket_expired_or_invalid')
    if ticket.owner != owner(request, principal):
        fail(404, 'upload_not_found')
    try:
        result = existing_result(request, ticket)
        if result is None:
            admission = ticket.admission
            with TemporaryDirectory() as directory:
                source = Path(directory) / 'staged'
                download_stage(config, ticket, source)
                if admission.purpose == 'curriculum':
                    from app.modules.curriculum.router import importer, import_error
                    from app.modules.curriculum.import_service import ImportFailure
                    from app.modules.curriculum.workbook_parser import WorkbookRejected
                    try:
                        result = importer(request).upload(owner=ticket.owner, curriculum_id=admission.curriculum_id,
                            filename=admission.filename, payload=source.read_bytes(), reserved_id=ticket.id)
                    except (ImportFailure, WorkbookRejected) as error:
                        import_error(error)
                else:
                    with source.open('rb') as stream:
                        upload = UploadFile(stream, filename=admission.filename,
                                            headers=Headers({'content-type': admission.mime_type}))
                        if admission.purpose == 'resource':
                            from app.modules.resources.router import store_upload
                            result = store_upload(request, principal, upload, reserved_id=ticket.id)
                        else:
                            from app.modules.bulk_transfer.router import preview_file
                            result = preview_file(request, principal, admission.kind, upload, reserved_id=ticket.id)
    except IntegrityError:
        # Stable target IDs settle concurrent completions without session locks,
        # including when the database uses a transaction pooler.
        result = existing_result(request, ticket)
        if result is None:
            raise
    remove_stage(config, ticket)
    return result


@router.post('/export')
def export_artifact(request: Request, principal: Admin, body: Export):
    """Reuse owner-module export generation, then return only an expiring private URL."""
    config = request.app.state.settings
    origin, bucket, headers = cloud(config)
    owner(request, principal)
    if body.purpose == 'curriculum':
        from app.modules.curriculum.router import export_curriculum
        response = export_curriculum(request, body.curriculum_id, principal,
                                     source=body.source, revision_number=body.revision_number)
    else:
        from app.modules.bulk_transfer.router import export, template
        response = (template(request, principal, body.kind) if body.action == 'template'
                    else export(request, principal, body.kind, package=body.package))
    content = response.body
    if len(content) > 120 * 1024 * 1024:
        fail(413, 'export_too_large')
    mime = response.headers.get('content-type', 'application/octet-stream').split(';')[0]
    suffix = '.zip' if mime == 'application/zip' else '.xlsx'
    key = 'transfer-exports/' + uuid4().hex + suffix
    filename = re.search(r'filename="([^"]+)"', response.headers.get('content-disposition', ''))
    filename = filename.group(1) if filename else 'flame-export' + suffix
    try:
        with httpx.Client(timeout=120, follow_redirects=False) as client:
            uploaded = client.post(origin + '/object/' + bucket + '/' + key,
                headers={**headers, 'content-type': mime, 'x-upsert': 'false'}, content=content)
            if uploaded.status_code not in (200, 201):
                fail(503, 'direct_storage_unavailable')
            signed = client.post(origin + '/object/sign/' + bucket + '/' + key,
                                 headers=headers, json={'expiresIn': 300})
            if signed.status_code != 200:
                fail(503, 'direct_storage_unavailable')
            url = signed.json().get('signedURL')
            if not isinstance(url, str) or not url.startswith('/object/sign/' + bucket + '/' + key + '?'):
                fail(503, 'direct_storage_unavailable')
    except (httpx.HTTPError, ValueError, AttributeError):
        fail(503, 'direct_storage_unavailable')
    return {'url': origin + url, 'storage_origin': config.supabase_url,
            'filename': filename, 'expires_in': 300, 'byte_size': len(content)}
