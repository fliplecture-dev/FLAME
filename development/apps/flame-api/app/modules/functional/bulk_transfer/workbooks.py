"""Bounded workbook/package decoding and catalogue workbook export."""
import hashlib
import json
from uuid import uuid4
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.modules.functional.bulk_transfer.errors import fail
from app.modules.functional.bulk_transfer.template_guidance import resource_guidance

KINDS = ("resources", "teachers", "students", "activity-kits")
COMMON = ["record_key", "existing_id"]
EDUCATION = ["concepts", "synonyms", "grade_min", "grade_max", "teaching_languages", "purposes"]
FIELDS = {
    "resources": COMMON + ["title", "topic", "type", "external_url", "file", "description", "alt_text", "caption_or_transcript", "accessibility_note", "rights_basis", "rights_holder", "rights_reference", "attribution", "distribution_scope", "availability"] + EDUCATION,
    "teachers": COMMON + ["name", "bio", "qualifications", "experience", "languages", "availability"] + EDUCATION,
    "students": COMMON + ["display_name", "email"],
    "activity-kits": COMMON + ["name", "description", "image_files", "materials", "instructions", "supervision", "cost", "currency", "availability", "supplier_url"] + EDUCATION,
}
MAPPINGS = ["record_key", "topic_id", "subtopic_id", "subject_id"]
MAX_PACKAGE = 120 * 1024 * 1024
MAX_EXPANDED = 200 * 1024 * 1024


def assign_record_keys(kind, records):
    """Generate stable import aliases without changing explicit keys or catalogue IDs."""
    for record in records:
        if record.get("record_key") is not None and str(record["record_key"]).strip():
            continue
        existing = str(record.get("existing_id") or "").strip().lower()
        if existing:
            identity = ["existing", existing]
        elif kind == "resources":
            media_type = str(record.get("type") or "").strip().lower()
            media_type = {"diagram":"image", "animation":"video", "external_video":"video"}.get(media_type, media_type)
            url = str(record.get("external_url") or "").strip()
            identity = [media_type, "url", url] if url else [media_type, "title", str(record.get("title") or "").strip().casefold()]
        else:
            identity = [str(record.get("email" if kind == "students" else "name") or "").strip().casefold()]
        token = hashlib.sha256(json.dumps([kind, identity], ensure_ascii=False).encode("utf-8")).hexdigest()[:24] if identity[-1] else uuid4().hex
        record["record_key"] = "auto-" + kind + "-" + token


def decode(kind, filename, content):
    assets = {}
    if filename.lower().endswith(".zip"):
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 250 or sum(e.file_size for e in entries) > MAX_EXPANDED:
                    fail(413, "Package exceeds 250 entries or 200 MiB expanded.")
                names = set()
                for entry in entries:
                    name = entry.filename
                    path = PurePosixPath(name)
                    if ("\\" in name or path.is_absolute() or ".." in path.parts or ":" in name or
                            (entry.external_attr >> 16) & 0o170000 == 0o120000 or entry.flag_bits & 1 or name in names):
                        fail(422, "Unsafe, duplicate or encrypted ZIP entry.")
                    names.add(name)
                    if entry.is_dir():
                        continue
                    if name != kind + ".xlsx" and not name.startswith("assets/"):
                        fail(422, "ZIP must contain " + kind + ".xlsx at its root and media inside assets/.")
                    if entry.file_size > 100 * 1024 * 1024:
                        fail(413, "A ZIP entry exceeds 100 MiB.")
                if kind + ".xlsx" not in names:
                    fail(422, "Package is missing " + kind + ".xlsx at its root.")
                workbook_bytes = archive.read(kind + ".xlsx")
                assets = {e.filename: archive.read(e) for e in entries if not e.is_dir() and e.filename.startswith("assets/")}
        except BadZipFile:
            fail(422, "The ZIP package is corrupt.")
    elif filename.lower().endswith(".xlsx"):
        workbook_bytes = content
    else:
        fail(422, "Upload an .xlsx workbook or a .zip package.")
    if len(workbook_bytes) > 10 * 1024 * 1024:
        fail(413, "Workbook exceeds 10 MiB.")
    # Bound the XLSX container itself before XML decoding.
    try:
        with ZipFile(BytesIO(workbook_bytes)) as zipped:
            if len(zipped.infolist()) > 500 or sum(e.file_size for e in zipped.infolist()) > 40 * 1024 * 1024:
                fail(413, "Workbook expands beyond the supported limit.")
        book = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError):
        fail(422, "The Excel workbook is invalid.")
    try:
        if "Records" not in book.sheetnames:
            fail(422, "Workbook requires a Records sheet.")
        def read(name, allowed, maximum):
            if name not in book.sheetnames:
                return []
            sheet = book[name]
            if (sheet.max_row or 0) > maximum + 1 or (sheet.max_column or 0) > 60:
                fail(413, f"{name} exceeds {maximum} rows or 60 columns.")
            rows = iter(sheet.iter_rows())
            headers = [str(c.value or "").strip() for c in next(rows)]
            while headers and not headers[-1]: headers.pop()
            if len(headers) > 60: fail(413, "Too many workbook columns.")
            if len(set(headers)) != len(headers) or any(h not in allowed for h in headers):
                fail(422, "Unknown or duplicate columns in " + name + ". Download the template.")
            output = []
            for number, cells in enumerate(rows, 2):
                if number > maximum + 1: fail(413, f"{name} exceeds {maximum} rows.")
                if any(c.value is not None for c in cells[len(headers):]): fail(422, f"{name} row {number}: data has no column heading.")
                if any(c.data_type == "f" for c in cells):
                    fail(422, f"{name} row {number}: formulas are not accepted.")
                row = {h: c.value for h, c in zip(headers, cells) if c.value is not None}
                if not row:
                    continue
                if any(not isinstance(v, (str, int, float, bool)) for v in row.values()):
                    fail(422, f"{name} row {number}: unsupported cell value.")
                row["_row"] = number
                output.append(row)
            return output
        records = read("Records", FIELDS[kind] + ["published", "version", "status"], 2000)
        mappings = read("Mappings", MAPPINGS, 10000)
        assign_record_keys(kind, records)
        if any(not str(row.get("record_key") or "").strip() for row in mappings):
            fail(422, "Mappings rows require a record_key matching Records. Supply your own keys for Excel mappings, or leave Mappings empty and map after import.")
        return {"records": records, "mappings": mappings}, assets
    finally:
        book.close()


def workbook(kind, records=(), mappings=(), paths=(), *, instructions=None, template=False):
    book = Workbook()
    book.remove(book.active)
    headers = [field for field in FIELDS[kind] if field != "existing_id"] if template else FIELDS[kind] + ["published", "version", "status"]
    sheets = [("Records", headers, records)]
    if kind != "students":
        sheets += [("Mappings", MAPPINGS, mappings), ("Paths", ["topic_id", "subtopic_id", "subject_id", "label"], paths)]
    for name, headers, rows in sheets:
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(h) for h in headers])
            # Data is always text where supplied as text, even if it starts with '='.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "C2" if name == "Records" else "B2"
        if name == "Records" and kind == "resources" and template:
            resource_guidance(sheet)
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244C78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 34
        for col in sheet.columns:
            header = col[0].value
            sheet.column_dimensions[col[0].column_letter].width = 42 if header in {"description", "label", "instructions", "external_url"} else 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide = book.create_sheet("Guide")
    notes = instructions or ["Fill Records. record_key is optional: FLAME generates a stable key when blank or omitted. Templates do not require existing_id.",
        "Existing nonempty values are preserved. Empty values may be filled on drafts; valid mappings are added, never removed.",
        "Published resource media/details remain unchanged; only valid mappings can be added.",
        "Use | between list values. Use assets/relative-file-name for bundled files.",
        "rights_basis: owned, licensed or permission only when verified; leave blank for an unreviewed draft. unknown is not accepted.",
        "purposes is optional; allowed values: explain, demonstrate, practise, human_help. Separate multiple allowed values with |. Other verbs such as explore/discuss are not accepted.",
        "Resource topic is an optional grouping label next to title. Filter by Topic label after import; it does not create curriculum mappings.",
        "For Excel mappings, supply matching record_key values in Records and Mappings. Otherwise leave Mappings empty and map after import. Copy path IDs from Paths.",
        "Interactive rows may omit file/external_url to create an HTML pending draft; add HTML in Edit later. Preview/publication require media.",
        "published, version and status are informational; import never publishes or grants access.",
        "Students become pending profiles; verified sign-in is required to activate them. No passwords or invitations are imported."]
    guide.append(["FLAME bulk transfer — " + kind])
    for note in notes:
        guide.append([note])
    guide.column_dimensions["A"].width = 110
    guide.sheet_view.showGridLines = False
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 36
    out = BytesIO(); book.save(out)
    return out.getvalue()
