"""Round-trip Excel export for published, draft, and historical curriculum snapshots."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.modules.functional.curriculum.drafts import published_paths, published_revision
from app.modules.functional.curriculum.models import CurriculumDraft, CurriculumRevision

SCHEMA_VERSION = "flame-curriculum-1"
HEADERS = ("Curriculum Code", "Level", "Subject", "Domain", "Category",
           "Topic", "Subtopic", "Learning Outcome", "Source URL")


def export_workbook(session, curriculum, source: str, revision: int | None = None) -> bytes:
    if source == "draft":
        draft = session.get(CurriculumDraft, curriculum.id)
        paths = draft.paths if draft else published_paths(session, curriculum.id)
        exported_revision = "draft-" + str(draft.revision if draft else 0)
    elif source == "published":
        paths = published_paths(session, curriculum.id)
        exported_revision = str(published_revision(session, curriculum))
    elif source == "revision" and revision is not None:
        item = session.scalar(select(CurriculumRevision).where(
            CurriculumRevision.curriculum_id == curriculum.id,
            CurriculumRevision.number == revision))
        paths = (item.details or {}).get("snapshot") if item else None
        if paths is None:
            raise ValueError("revision_snapshot_unavailable")
        exported_revision = str(revision)
    else:
        raise ValueError("invalid_export_source")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Curriculum"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="204B7B")
    for row in paths:
        values = row["values"]
        sheet.append([
            curriculum.code, values["level"], values["subject"], values["domain"],
            values.get("category") or None, values["topic"], values["subtopic"],
            None, None,
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (22, 18, 24, 24, 24, 28, 34, 40, 35)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["FLAME curriculum round-trip workbook"])
    instructions.append(["Edit rows only on the Curriculum sheet."])
    instructions.append(["Level, Subject, Domain, Topic and Subtopic are required. Category is optional."])
    instructions.append(["Importing this workbook proposes a complete hierarchy replacement."])
    instructions.append(["Rows absent from the workbook are proposed for removal from the draft."])

    metadata = workbook.create_sheet("Metadata")
    for key, value in (
        ("schema_version", SCHEMA_VERSION),
        ("curriculum_id", str(curriculum.id)),
        ("curriculum_code", curriculum.code),
        ("curriculum_name", curriculum.name),
        ("source", source),
        ("exported_revision", exported_revision),
        ("exported_at", datetime.now(timezone.utc).isoformat()),
    ):
        metadata.append([key, value])

    identifiers = workbook.create_sheet("Stable IDs")
    identifiers.append([*("Level ID", "Subject ID", "Domain ID", "Category ID", "Topic ID", "Subtopic ID"),
                        *HEADERS[1:7]])
    for row in paths:
        ids, values = row.get("ids") or {}, row["values"]
        identifiers.append([
            ids.get("level"), ids.get("subject"), ids.get("domain"), ids.get("category"),
            ids.get("topic"), ids.get("subtopic"),
            values["level"], values["subject"], values["domain"], values.get("category") or None,
            values["topic"], values["subtopic"],
        ])
    identifiers.sheet_state = "hidden"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()

