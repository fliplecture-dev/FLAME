"""Real upload + resource workflow, using private temp media and rollback-only DB data."""
from io import BytesIO
from uuid import UUID, uuid4
from PIL import Image
from openpyxl import load_workbook
from sqlalchemy import select, func
from pydantic import SecretStr
from app.modules.functional.curriculum.models import Subtopic
from app.modules.functional.resources.models import Resource, ResourceTopic, ResourceUpload
from test_curriculum_workspace import isolated_client


def test_resource_upload_retries_paths_validation_and_publication(tmp_path):
    with isolated_client() as client:
        config = client.app.state.settings
        config.resource_storage_backend = 'local'; config.resource_local_dir = str(tmp_path)
        config.bff_shared_key = SecretStr('isolated-resource-test-signing-key')
        code = 'media-' + uuid4().hex[:10]
        c = client.post('/api/v1/admin/curricula', json={'code': code, 'name': 'Media fixture'}).json()
        cid = c['id']; root = '/api/v1/admin/curricula/' + cid; resources = '/api/v1/admin/resources'
        def child(kind, parent, name):
            r = client.post('/api/v1/admin/curriculum/' + kind, json={'parent_id': parent, 'name': name})
            assert r.status_code == 201, r.text
            return r.json()['id']
        level = child('level', cid, 'Grade 1'); subject = child('subject', level, 'Science')
        domain = child('domain', subject, 'Motion'); topic = child('topic', domain, 'Gravity')
        leaf = child('subtopic', topic, 'Falling'); other = child('topic', domain, 'Orbit')
        other_leaf = child('subtopic', other, 'Moon')
        with client.app.state.database.session() as session, session.begin():
            session.add_all([Subtopic(topic_id=UUID(topic), name=f'Extra {i}', normalized_name=f'extra {i}') for i in range(103)])
        def publish_curriculum():
            t = client.get(root + '/tree').json()
            r = client.post('/api/v1/admin/curriculum/curriculum/' + cid + '/publish', json={
                'expected_version': t['version'], 'expected_draft_revision': t['draft_revision']})
            assert r.status_code == 200, r.text
            return r.json()
        publish_curriculum()
        choices = client.get(resources + '/paths', params={'curriculum_id': cid})
        assert choices.status_code == 200, choices.text
        assert len(choices.json()['items']) > 100
        image = BytesIO(); Image.new('RGB', (3, 3), 'blue').save(image, format='PNG')
        uploaded = client.post(resources + '/uploads', files={'file': ('isolated.png', image.getvalue(), 'image/png')})
        assert uploaded.status_code == 201, uploaded.text
        upload_id = uploaded.json()['upload_id']; creation_id = str(uuid4())
        paths = [{'topic_id': topic, 'subtopic_id': None}, {'topic_id': topic, 'subtopic_id': leaf},
                 {'topic_id': other, 'subtopic_id': None}, {'topic_id': other, 'subtopic_id': other_leaf}]
        payload = {'creation_id': creation_id, 'title': 'Isolated Gravity', 'type': 'image', 'upload_id': upload_id,
                   'rights_basis': 'owned', 'distribution_scope': 'members', 'paths': paths}
        bad = client.post(resources, json={**payload, 'paths': [*paths, {'topic_id': other, 'subtopic_id': leaf}]})
        assert bad.status_code == 422 and bad.json()['error']['code'] == 'resource_paths_unavailable'
        assert bad.json()['error']['invalid_paths'] == [{'topic_id': other, 'subtopic_id': leaf}]
        assert bad.json()['request_id'] == bad.headers['x-request-id']
        with client.app.state.database.session() as session:
            assert session.get(ResourceUpload, UUID(upload_id)).resource_id is None
            assert session.get(Resource, UUID(creation_id)) is None
        # Retry the atomic resource + association stage using the same uploaded file.
        saved = client.post(resources, json=payload)
        assert saved.status_code == 201, saved.text
        row = saved.json(); rid = row['id']; assert rid == creation_id
        assert row['source'] == 'upload' and row['rights_basis'] == 'owned' and row['distribution_scope'] == 'members'
        assert len(row['paths']) == len(row['path_details']) == 4
        repeated = client.post(resources, json=payload)
        assert repeated.status_code == 409 and repeated.json()['error']['code'] == 'resource_creation_exists'
        assert client.get(resources + '/' + creation_id).json()['version'] == row['version']
        with client.app.state.database.session() as session:
            assert session.scalar(select(func.count()).select_from(ResourceUpload).where(ResourceUpload.id == UUID(upload_id))) == 1
        preview = client.get(resources + '/' + rid + '/preview')
        assert preview.status_code == 200, preview.text
        assert client.get('/api/v1/admin/resources/' + rid + '/preview-content').content == image.getvalue()
        denied = client.post(resources + '/' + rid + '/publish', json={'expected_version': row['version']})
        assert denied.status_code == 422 and denied.json()['error']['code'] == 'distribution_rights_required'
        assert denied.json()['error']['fields'] == ['rights_holder']
        assert client.get(resources + '/' + rid).json()['published'] is False
        assert client.get('/api/v1/resources', params={'topic_id': topic}).json()['items'] == []
        assert client.patch(resources + '/' + rid, json={'expected_version': 99, 'title': 'Stale'}).json()['error']['code'] == 'stale_version'
        all_paths = [{'topic_id': item['topic_id'], 'subtopic_id': item['subtopic_id']} for item in choices.json()['items']]
        updated = client.patch(resources + '/' + rid, json={'expected_version': row['version'], 'rights_holder': 'Fixture owner', 'paths': all_paths})
        assert updated.status_code == 200, updated.text
        row = updated.json(); assert len(row['paths']) > 100
        blocked = client.post(resources + '/' + rid + '/publish', json={'expected_version': row['version']})
        assert blocked.status_code == 422 and blocked.json()['error']['code'] == 'alt_text_required'
        updated = client.patch(resources + '/' + rid, json={'expected_version': row['version'], 'alt_text': 'A blue gravity illustration', 'paths': paths})
        assert updated.status_code == 200, updated.text
        row = updated.json()
        failed_edit = client.patch(resources + '/' + rid, json={'expected_version': row['version'], 'title': 'Must roll back',
            'paths': [{'topic_id': topic, 'subtopic_id': str(uuid4())}]})
        assert failed_edit.status_code == 422
        assert client.get(resources + '/' + rid).json()['title'] == 'Isolated Gravity'
        assert client.get(resources + '/' + rid).json()['version'] == row['version']
        pub = client.post(resources + '/' + rid + '/publish', json={'expected_version': row['version']})
        assert pub.status_code == 200, pub.text
        for path in paths:
            params = {k: v for k, v in path.items() if v is not None}
            found = client.get('/api/v1/resources', params=params)
            assert found.status_code == 200 and any(item['id'] == rid for item in found.json()['items']), found.text
        delivery = client.get('/api/v1/resources/' + rid + '/delivery')
        assert delivery.status_code == 200 and client.get(delivery.json()['url']).content == image.getvalue()
        # Actual replacement removes a path; resource IDs/associations remain intact.
        exported = client.get(root + '/export', params={'source': 'published'}).content
        wb = load_workbook(BytesIO(exported)); sheet = wb['Curriculum']
        headings = [cell.value for cell in sheet[1]]; column = headings.index('Subtopic') + 1
        for number in range(sheet.max_row, 1, -1):
            if sheet.cell(number, column).value == 'Falling': sheet.delete_rows(number)
        out = BytesIO(); wb.save(out)
        imports = '/api/v1/admin/curriculum/imports'
        batch = client.post(imports, params={'curriculum_id': cid}, files={'file': ('removed.xlsx', out.getvalue())}).json()['import_id']
        preview = client.post(imports + '/' + batch + '/preview', json={'sheet_levels': {'Curriculum': '__ALL_LEVELS__'}})
        assert preview.status_code == 200, preview.text
        applied = client.post(imports + '/' + batch + '/commit', headers={'idempotency-key': uuid4().hex}, json={
            'preview_revision': preview.json()['preview_revision'], 'confirm': 'REPLACE ' + code})
        assert applied.status_code == 200, applied.text
        publish_curriculum()
        row = client.get(resources + '/' + rid).json()
        assert any(detail['valid'] is False for detail in row['path_details'])
        stale_path = client.patch(resources + '/' + rid, json={'expected_version': row['version'], 'paths': paths})
        assert stale_path.status_code == 422 and stale_path.json()['error']['code'] == 'resource_paths_unavailable'
        assert len(client.get(resources + '/' + rid).json()['paths']) == 4
        repaired = client.patch(resources + '/' + rid, json={'expected_version': row['version'], 'paths': [p for p in paths if p['subtopic_id'] != leaf]})
        assert repaired.status_code == 200, repaired.text
        assert len(repaired.json()['paths']) == 3
