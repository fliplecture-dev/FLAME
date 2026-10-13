"""Rollback-only transfers: no provider calls, real accounts or permanent catalogue changes."""
from io import BytesIO
from uuid import uuid4, UUID
from zipfile import ZipFile
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import pytest
from PIL import Image
from openpyxl import load_workbook
from app.modules.bulk_transfer.workbooks import workbook, decode
from test_curriculum_workspace import isolated_client


def post_preview(client, kind, rows, links=(), files=None):
    data = workbook(kind, rows, links)
    if files is not None:
        out=BytesIO()
        with ZipFile(out, 'w') as z:
            z.writestr(kind+'.xlsx', data)
            for name, content in files.items(): z.writestr(name, content)
        data=out.getvalue(); name=kind+'.zip'
    else: name=kind+'.xlsx'
    r=client.post('/api/v1/admin/bulk/'+kind+'/preview', files={'file':(name,data)})
    assert r.status_code==200, r.text
    return r.json()


def commit(client, job):
    return client.post('/api/v1/admin/bulk/imports/'+job['id']+'/commit', json={'preview_token':job['preview_token'],'confirm':True})


def path_fixture(client, published=False):
    c=client.post('/api/v1/admin/curricula', json={'code':'bulk-'+uuid4().hex[:8],'name':'Bulk fixture'}).json()
    def child(kind,parent,name):
        r=client.post('/api/v1/admin/curriculum/'+kind,json={'parent_id':parent,'name':name});assert r.status_code==201,r.text;return r.json()['id']
    level=child('level',c['id'],'Grade 1');subject=child('subject',level,'Science');domain=child('domain',subject,'Motion');topic=child('topic',domain,'Gravity');leaf=child('subtopic',topic,'Falling objects')
    if published:
        tree=client.get('/api/v1/admin/curricula/'+c['id']+'/tree').json()
        result=client.post('/api/v1/admin/curriculum/curriculum/'+c['id']+'/publish',json={'expected_version':tree['version'],'expected_draft_revision':tree['draft_revision']})
        assert result.status_code==200,result.text
    return topic,leaf,subject


def test_four_catalogues_add_merge_export_and_retry(monkeypatch):
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        topic,leaf,subject=path_fixture(client)
        key=uuid4().hex
        png=BytesIO();Image.new('RGB',(8,8),'blue').save(png,format='PNG')
        media=[{'record_key':key+'-media','title':'Bulk media '+key,'type':'image','file':'assets/example.png','alt_text':'A blue square fixture.', 'rights_basis':'owned','rights_holder':'Fixture owner','distribution_scope':'members','concepts':'gravity'}]
        mappings=[{'record_key':media[0]['record_key'],'topic_id':topic}]
        job=post_preview(client,'resources',media,mappings,{'assets/example.png':png.getvalue()})
        assert job['counts']=={'add':1,'merge':0,'unchanged':0,'invalid':0}
        saved=commit(client,job);assert saved.status_code==200,saved.text
        rid=saved.json()['result']['ids'][media[0]['record_key']]
        assert commit(client,job).json()['result']==saved.json()['result']
        row=client.get('/api/v1/admin/resources/'+rid).json()
        assert not row['published'] and row['learning_metadata']['concepts']==['gravity']
        later=[{'record_key':media[0]['record_key'], 'title':'Keep original title', 'description':'New missing description','concepts':'force'}]
        more=mappings+[{'record_key':media[0]['record_key'],'topic_id':topic,'subtopic_id':leaf}]
        job=post_preview(client,'resources',later,more)
        assert job['counts']['merge']==1 and 'title' in job['rows'][0]['preserved_fields']
        assert commit(client,job).status_code==200
        changed=client.get('/api/v1/admin/resources/'+rid).json()
        assert changed['title']==row['title'] and changed['description']=='New missing description' and changed['upload_id']==row['upload_id'] and len(changed['paths'])==2
        assert changed['learning_metadata']['concepts']==['gravity','force']
        teacher=[{'record_key':key+'-teacher','name':'Bulk teacher '+key,'languages':'English','availability':'unknown'}]
        links=[{'record_key':teacher[0]['record_key'],'subject_id':subject,'topic_id':topic}]
        first=commit(client,post_preview(client,'teachers',teacher,links));assert first.status_code==200,first.text
        tid=first.json()['result']['ids'][teacher[0]['record_key']]
        merged=commit(client,post_preview(client,'teachers',[{**teacher[0],'name':'Preserved name','bio':'Fill missing bio','languages':'Tamil'}],links));assert merged.status_code==200,merged.text
        t=client.get('/api/v1/admin/teachers/'+tid).json();assert t['name']==teacher[0]['name'] and t['languages']==['English','Tamil'] and not t['published']
        kits=[{'record_key':key+'-kit','name':'Bulk kit '+key,'materials':'Ball|Tray','instructions':'Observe a falling ball.','supervision':'Adult supervision','image_files':'assets/kit.png'}]
        klinks=[{'record_key':kits[0]['record_key'],'topic_id':topic,'subtopic_id':leaf}]
        first=commit(client,post_preview(client,'activity-kits',kits,klinks,{'assets/kit.png':png.getvalue()}));assert first.status_code==200,first.text
        kid=first.json()['result']['ids'][kits[0]['record_key']]
        changed=commit(client,post_preview(client,'activity-kits',[{**kits[0],'materials':'Ball|Cup','description':'A missing description'}],klinks));assert changed.status_code==200,changed.text
        k=client.get('/api/v1/admin/activity-kits/'+kid).json();assert k['materials']==['Ball','Tray','Cup'] and len(k['image_upload_ids'])==1 and not k['published']
        students=[{'record_key':key+'-student','display_name':'Sample student','email':key+'@example.test'}]
        first=commit(client,post_preview(client,'students',students));assert first.status_code==200,first.text
        sid=first.json()['result']['ids'][students[0]['record_key']]
        same=commit(client,post_preview(client,'students',[{**students[0],'display_name':'Do not overwrite'}]));assert same.status_code==200,same.text
        s=client.get('/api/v1/admin/students/'+sid).json();assert s['display_name']=='Sample student' and s['status']=='pending'
        for kind, ident in [('resources',rid),('teachers',tid),('activity-kits',kid),('students',sid)]:
            export=client.get('/api/v1/admin/bulk/'+kind+'/export');assert export.status_code==200,export.text[:100]
            book=load_workbook(BytesIO(export.content),read_only=True);rows=list(book['Records'].values);headers=rows[0]
            assert any(dict(zip(headers,r))['existing_id']==ident for r in rows[1:]);book.close()
            template=client.get('/api/v1/admin/bulk/'+kind+'/template');assert template.status_code==200
        # Isolate package export from existing real catalogue assets.
        import app.modules.bulk_transfer.router as router
        original=router.export_data
        def only_fixture(db,kind):
            records,maps,assets=original(db,kind)
            return [r for r in records if r['existing_id']==rid], [m for m in maps if m['record_key']==media[0]['record_key']], {n:v for n,v in assets.items() if rid in n}
        monkeypatch.setattr(router,'export_data',only_fixture)
        bundle=client.get('/api/v1/admin/bulk/resources/export?package=true');assert bundle.status_code==200,bundle.text[:100] if not bundle.status_code==200 else ''
        with ZipFile(BytesIO(bundle.content)) as z: assert z.read('assets/'+rid+'.png')==png.getvalue()


def test_stale_preview_invalid_rows_and_cancellation():
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        key=uuid4().hex
        teacher={'record_key':key,'name':'Stale teacher '+key}
        saved=commit(client,post_preview(client,'teachers',[teacher]));tid=saved.json()['result']['ids'][key]
        job=post_preview(client,'teachers',[{**teacher,'bio':'Missing bio'}])
        current=client.get('/api/v1/admin/teachers/'+tid).json()
        assert client.patch('/api/v1/admin/teachers/'+tid,json={'expected_version':current['version'],'bio':'Other tab saved'}).status_code==200
        stale=commit(client,job);assert stale.status_code==409 and 'Nothing was imported' in stale.json()['error']['message']
        latest=client.post('/api/v1/admin/bulk/imports/'+job['id']+'/review',json={}).json()
        assert latest['counts']['unchanged']==1
        assert commit(client,latest).status_code==200
        assert client.get('/api/v1/admin/teachers/'+tid).json()['bio']=='Other tab saved'
        invalid=post_preview(client,'students',[{'record_key':key+'a','email':'not-an-email'}, {'record_key':key+'b','email':key+'@example.test'}])
        assert invalid['counts']['invalid']==1 and commit(client,invalid).status_code==422
        assert client.post('/api/v1/admin/bulk/imports/'+invalid['id']+'/cancel',json={}).status_code==200
        assert commit(client,invalid).status_code==409


def test_in_file_duplicates_merge_and_unsafe_packages():
    with isolated_client() as client:
        key=uuid4().hex
        rows=[{'record_key':key+'a','name':'Duplicate '+key}, {'record_key':key+'b','name':'Duplicate '+key,'bio':'Additional bio'}, {'record_key':key+'a','name':'Do not change the first row name','qualifications':'Additional qualification'}]
        saved=commit(client,post_preview(client,'teachers',rows));assert saved.status_code==200,saved.text
        ids=saved.json()['result']['ids'];assert ids[key+'a']==ids[key+'b']
        profile=client.get('/api/v1/admin/teachers/'+ids[key+'a']).json();assert profile['bio']=='Additional bio' and profile['name']=='Duplicate '+key and profile['qualifications']=='Additional qualification'
        out=BytesIO()
        with ZipFile(out,'w') as z: z.writestr('../escape.xlsx',b'x')
        with pytest.raises(Exception) as error: decode('teachers','teachers.zip',out.getvalue())
        assert getattr(error.value,'status_code',None)==422


def test_imported_student_requires_verified_provider_identity(monkeypatch):
    with isolated_client() as client:
        key=uuid4().hex; email=key+'@example.test'; sub=str(uuid4())
        result=commit(client,post_preview(client,'students',[{'record_key':key,'email':email,'display_name':'Imported name'}]));sid=result.json()['result']['ids'][key]
        import app.modules.identity.router as auth
        monkeypatch.setattr(auth,'_session',lambda request,raw:({'sub':sub,'email':email},'opaque-token'))
        request=SimpleNamespace(app=client.app)
        monkeypatch.setattr(auth,'_provider',lambda *args,**kwargs:{'id':sub,'email':email,'email_confirmed_at':None})
        with pytest.raises(Exception): auth._principal(request,'session')
        monkeypatch.setattr(auth,'_provider',lambda *args,**kwargs:{'id':sub,'email':email,'email_confirmed_at':'verified'})
        principal=auth._principal(request,'session')
        assert principal['id']==sid and principal['role']=='member' and principal['display_name']=='Imported name'


def test_published_resource_merge_keeps_media_and_publication():
    with isolated_client() as client:
        client.app.state.settings.resource_allowed_hosts='media.example.test'
        topic,leaf,_=path_fixture(client,published=True); key=uuid4().hex
        rows=[{'record_key':key,'title':'Published bulk '+key,'type':'image','external_url':'https://media.example.test/'+key+'.png','alt_text':'A falling ball diagram.','rights_basis':'owned','rights_holder':'Fixture','distribution_scope':'members','concepts':'gravity'}]
        links=[{'record_key':key,'topic_id':topic}]
        saved=commit(client,post_preview(client,'resources',rows,links));assert saved.status_code==200,saved.text
        rid=saved.json()['result']['ids'][key]
        current=client.get('/api/v1/admin/resources/'+rid).json()
        published=client.post('/api/v1/admin/resources/'+rid+'/publish',json={'expected_version':current['version']});assert published.status_code==200,published.text
        updates=[{**rows[0],'description':'Do not change published details','concepts':'force'}]
        links.append({'record_key':key,'topic_id':topic,'subtopic_id':leaf})
        saved=commit(client,post_preview(client,'resources',updates,links));assert saved.status_code==200,saved.text
        row=client.get('/api/v1/admin/resources/'+rid).json()
        assert row['published'] and row['description'] is None and row['learning_metadata']['concepts']==['gravity'] and len(row['paths'])==2


def test_package_cancel_removes_only_staged_files_and_owner_isolation():
    from app.modules.identity.dependencies import require_admin
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        key=uuid4().hex;png=BytesIO();Image.new('RGB',(4,4),'blue').save(png,format='PNG')
        rows=[{'record_key':key,'title':'Cancelled '+key,'type':'animation','file':'assets/sample.png'}]
        job=post_preview(client,'resources',rows,files={'assets/sample.png':png.getvalue()})
        assert job['counts']['invalid']==1  # Actual image type is incompatible with Animation.
        from pathlib import Path
        assert len(list(Path(directory).rglob('*.png')))==1
        assert client.post('/api/v1/admin/bulk/imports/'+job['id']+'/cancel',json={}).status_code==200
        assert list(Path(directory).rglob('*.png'))==[]
        owner_job=post_preview(client,'students',[{'record_key':key,'email':key+'@example.test'}])
        from app.modules.identity.models import StudentProfile, AdminRole
        other=uuid4()
        with client.app.state.database.session() as db, db.begin():
            db.add_all([StudentProfile(id=other,auth_subject=str(other)),AdminRole(auth_subject=str(other))])
        client.app.dependency_overrides[require_admin]=lambda:{'id':str(other),'role':'admin'}
        assert client.get('/api/v1/admin/bulk/imports/'+owner_job['id']).status_code==404
        assert commit(client,owner_job).status_code==404


def test_generated_samples_decode_and_formulas_are_rejected():
    from pathlib import Path
    root=Path(__file__).resolve().parents[4]/'project-stages/bulk-import-samples'
    for kind in ('resources','teachers','students','activity-kits'):
        content=(root/kind/(kind+'.xlsx')).read_bytes()
        assert decode(kind,kind+'.xlsx',content)[0]['records']
    from openpyxl import Workbook
    book=Workbook();book.active.title='Records';book.active.append(['record_key','display_name','email']);book.active.append(['bad','=1+1','safe@example.test']);out=BytesIO();book.save(out)
    with pytest.raises(Exception) as error: decode('students','students.xlsx',out.getvalue())
    assert getattr(error.value,'status_code',None)==422


@pytest.mark.parametrize('kind,field,limit', [('resources','learning_metadata',20),('teachers','languages',20),('activity-kits','materials',100)])
def test_additive_merges_keep_existing_values_and_respect_limits(kind,field,limit):
    from app.modules.bulk_transfer.service import changes
    existing=[str(i) for i in range(limit)]
    before={field:{'concepts':existing} if field=='learning_metadata' else existing}
    body={field:{'concepts':['new']} if field=='learning_metadata' else ['new']}
    with pytest.raises(ValueError): changes(kind,before,body)
    assert (before[field]['concepts'] if field=='learning_metadata' else before[field])==existing


def test_external_video_fresh_inspection_timestamp_is_not_a_conflict(monkeypatch):
    import app.modules.bulk_transfer.service as service
    from app.modules.resources.external_video import parse_video
    calls=[]
    def inspected(url, config):
        calls.append(url)
        return {**parse_video(url),'title':'Provider fixture','description':'','thumbnail_url':None,
            'verification':'metadata_only','verified_at':'inspection-'+str(len(calls)),
            'student_playback_allowed':True,'playback_reason':'Standard player'}
    monkeypatch.setattr(service,'inspect_video',inspected)
    with isolated_client() as client:
        key=uuid4().hex
        job=post_preview(client,'resources',[{'record_key':key,'title':'External fixture '+key,'type':'external_video','external_url':'https://youtu.be/M7lc1UVf-VE'}])
        result=commit(client,job)
        assert result.status_code==200,result.text
        assert len(calls)>=3  # Actual validation is repeated; only its volatile timestamp is excluded.
        assert result.json()['result']['added']==1


@pytest.mark.parametrize("kind,record", [
    ("resources", {"title":"Automatic image", "type":"image", "file":"assets/photo.png"}),
    ("teachers", {"name":"Automatic teacher"}),
    ("students", {"display_name":"Automatic student", "email":"auto@example.test"}),
    ("activity-kits", {"name":"Automatic kit"}),
])
def test_automatic_record_keys_and_simple_templates(kind, record):
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title="Records"
    sheet.append(list(record));sheet.append(list(record.values()))
    data=BytesIO();book.save(data)
    first,_=decode(kind,kind+".xlsx",data.getvalue())
    second,_=decode(kind,kind+".xlsx",data.getvalue())
    assert first["records"][0]["record_key"].startswith("auto-"+kind+"-")
    assert first["records"][0]["record_key"]==second["records"][0]["record_key"]
    blank,_=decode(kind,kind+".xlsx",workbook(kind,[{**record,"record_key":""}]))
    assert blank["records"][0]["record_key"]==first["records"][0]["record_key"]
    explicit,_=decode(kind,kind+".xlsx",workbook(kind,[{**record,"record_key":"my-key"}]))
    assert explicit["records"][0]["record_key"]=="my-key"
    template=load_workbook(BytesIO(workbook(kind,template=True)))
    headers=[cell.value for cell in template["Records"][1]]
    assert "existing_id" not in headers and "record_key" in headers
    assert "existing_id" in [cell.value for cell in load_workbook(BytesIO(workbook(kind)))["Records"][1]]


def test_mapping_without_record_key_is_actionable():
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as caught:
        decode("resources","resources.xlsx",workbook("resources",[{"title":"Automatic","type":"image","file":"assets/x.png"}], [{"topic_id":str(uuid4())}]))
    assert "map after import" in caught.value.detail


def test_automatic_media_key_reimport_review_and_export():
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        title="Automatic media "+uuid4().hex
        png=BytesIO();Image.new("RGB",(8,8),"blue").save(png,format="PNG")
        records=[{"title":title,"type":"image","file":"assets/auto.png","alt_text":"A blue square"}]
        job=post_preview(client,"resources",records,files={"assets/auto.png":png.getvalue()})
        key=job["rows"][0]["record_key"];assert key.startswith("auto-resources-")
        refreshed=client.post("/api/v1/admin/bulk/imports/"+job["id"]+"/review",json={}).json()
        assert refreshed["rows"][0]["record_key"]==key
        receipt=commit(client,refreshed);assert receipt.status_code==200,receipt.text
        rid=receipt.json()["result"]["ids"][key]
        second=post_preview(client,"resources",records,files={"assets/auto.png":png.getvalue()})
        assert second["rows"][0]["record_key"]==key and second["rows"][0]["existing_id"]==rid
        assert second["counts"]["unchanged"]==1 and commit(client,second).status_code==200
        exported=client.get("/api/v1/admin/bulk/resources/export")
        data,_=decode("resources","resources.xlsx",exported.content)
        row=next(row for row in data["records"] if row.get("existing_id")==rid)
        assert row["record_key"]==key



def test_resource_topic_labels_import_filter_export_and_mapping_separation():
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        name="Topic-labelled media "+uuid4().hex
        png=BytesIO();Image.new("RGB",(8,8),"blue").save(png,format="PNG")
        records=[{"title":name,"topic":"Gravity | Force","type":"image","file":"assets/topics.png"}]
        job=post_preview(client,"resources",records,files={"assets/topics.png":png.getvalue()})
        assert job["counts"]["invalid"]==0
        saved=commit(client,job);assert saved.status_code==200,saved.text
        rid=saved.json()["result"]["ids"][job["rows"][0]["record_key"]]
        row=client.get("/api/v1/admin/resources/"+rid).json()
        assert row["learning_metadata"]["topic_labels"]==["Gravity","Force"] and row["paths"]==[]
        filtered=client.get("/api/v1/admin/resources",params={"q":name,"topic_label":"gravity"}).json()
        assert filtered["total"]==1 and filtered["items"][0]["id"]==rid
        assert client.get("/api/v1/admin/resources",params={"q":name,"topic_label":"Light"}).json()["total"]==0
        assert client.get("/api/v1/admin/resources",params={"q":name,"topic_label":"%"}).json()["total"]==0
        data,_=decode("resources","resources.xlsx",client.get("/api/v1/admin/bulk/resources/export").content)
        exported=next(record for record in data["records"] if record.get("existing_id")==rid)
        import json
        assert json.loads(exported["topic"])==["Gravity","Force"]
        template=load_workbook(BytesIO(client.get("/api/v1/admin/bulk/resources/template").content))
        headers=[cell.value for cell in template["Records"][1]]
        assert headers.index("topic")==headers.index("title")+1



def test_mixed_export_twenty_keyless_interactives_partial_reimport_and_errors():
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        prefix=uuid4().hex
        png=BytesIO();Image.new("RGB",(8,8),"blue").save(png,format="PNG")
        initial=[{"record_key":prefix+"-old", "title":"Old "+prefix,"type":"image","file":"assets/old.png","description":"Keep saved description"}]
        original=commit(client,post_preview(client,"resources",initial,files={"assets/old.png":png.getvalue()}))
        assert original.status_code==200,original.text
        old_id=original.json()["result"]["ids"][prefix+"-old"]
        original_row=client.get("/api/v1/admin/resources/"+old_id).json()
        exported,_=decode("resources","resources.xlsx",client.get("/api/v1/admin/bulk/resources/export").content)
        old=next(record for record in exported["records"] if record.get("existing_id")==old_id)
        additions=[{"title":"HTML "+prefix+" "+str(i),"type":"interactive","topic":"Gravity","file":"assets/new-"+str(i)+".html"} for i in range(20)]
        html=b'<!doctype html><html><head><title>Fixture</title></head><body><h1>Gravity</h1></body></html>'
        assets={row["file"]:html for row in additions}
        # Missing new assets yields invalid rows but a consistent renderable response.
        invalid=post_preview(client,"resources",[old,*additions])
        assert invalid["counts"]["invalid"]==20
        for row in invalid["rows"]:
            assert isinstance(row["errors"],list) and isinstance(row["preserved_fields"],list) and isinstance(row["fields_to_fill"],list)
        assert "upload a ZIP" in invalid["rows"][1]["errors"][0]
        assert commit(client,invalid).status_code==422
        # Existing exported media needs no bundled file; only new files are supplied.
        mixed=post_preview(client,"resources",[old,*additions],files=assets)
        assert mixed["counts"]=={"add":20,"merge":0,"unchanged":1,"invalid":0}
        assert mixed["rows"][0]["record_key"]==prefix+"-old"
        assert all(row["record_key"].startswith("auto-resources-") for row in mixed["rows"][1:])
        receipt=commit(client,mixed);assert receipt.status_code==200,receipt.text
        assert receipt.json()["result"]["added"]==20
        assert client.get("/api/v1/admin/resources/"+old_id).json()["upload_id"]==original_row["upload_id"]
        assert commit(client,mixed).json()["result"]==receipt.json()["result"]
        # Full repeated ZIP is unchanged; partial Interactive-only upload never removes omitted rows.
        repeated=post_preview(client,"resources",[old,*additions],files=assets)
        assert repeated["counts"]["unchanged"]==21 and repeated["counts"]["invalid"]==0
        partial=post_preview(client,"resources",additions[:2],files={row["file"]:html for row in additions[:2]})
        assert partial["counts"]["unchanged"]==2 and commit(client,partial).status_code==200
        for rid in receipt.json()["result"]["ids"].values():
            assert client.get("/api/v1/admin/resources/"+rid).status_code==200
        old["description"]="Do not overwrite";old["alt_text"]="Fill empty accessibility"
        merge=post_preview(client,"resources",[old])
        assert merge["counts"]["merge"]==1 and commit(client,merge).status_code==200
        saved=client.get("/api/v1/admin/resources/"+old_id).json()
        assert saved["description"]=="Keep saved description" and saved["alt_text"]=="Fill empty accessibility"



def test_resource_authoring_errors_are_reported_together():
    from app.modules.bulk_transfer.service import raw_body
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as caught:
        raw_body("resources", {"record_key":"new","title":"Video fixture","type":"video","purposes":"explore|discuss","rights_basis":"unknown"}, [], {})
    message=caught.value.detail
    assert "rights_basis:" in message and "purposes:" in message and "Media source:" in message
    valid=raw_body("resources", {"record_key":"new","title":"HTML fixture","type":"interactive","file":"assets/lab.html"}, [], {})
    assert valid["rights_basis"] is None and valid["learning_metadata"]["purposes"]==[]



def test_pending_interactive_import_then_html_upload_publish():
    with isolated_client() as client, TemporaryDirectory() as directory:
        client.app.state.settings.resource_local_dir=directory
        for media_type in ["image","video"]:
            invalid=client.post("/api/v1/admin/resources",json={"title":"No media","type":media_type})
            assert invalid.status_code==422
        name="Pending HTML "+uuid4().hex
        record={"title":name,"type":"interactive","rights_basis":"owned","rights_holder":"Test owner","distribution_scope":"members","accessibility_note":"Tab and Enter operate the controls."}
        job=post_preview(client,"resources",[record])
        assert job["counts"]["add"]==1 and job["counts"]["invalid"]==0
        result=commit(client,job);assert result.status_code==200,result.text
        rid=result.json()["result"]["ids"][job["rows"][0]["record_key"]]
        row=client.get("/api/v1/admin/resources/"+rid).json()
        assert row["source"]=="pending" and not row["media_ready"] and not row["published"]
        preview=client.get("/api/v1/admin/resources/"+rid+"/preview")
        assert preview.status_code==200 and preview.json()["launch_mode"]=="interactive_pending"
        blocked=client.post("/api/v1/admin/resources/"+rid+"/publish",json={"expected_version":row["version"]})
        assert blocked.status_code==422 and blocked.json()["error"]["code"]=="interactive_media_required"
        again=post_preview(client,"resources",[record]);assert again["counts"]["unchanged"]==1
        html=b'<!doctype html><html><head><title>Lab</title></head><body><h1>Gravity lab</h1></body></html>'
        upload=client.post("/api/v1/admin/resources/uploads",files={"file":("lab.html",html,"text/html")});assert upload.status_code==201,upload.text
        changed=client.patch("/api/v1/admin/resources/"+rid,json={"expected_version":row["version"],"upload_id":upload.json()["upload_id"]})
        assert changed.status_code==200,changed.text
        saved=changed.json();assert saved["media_ready"] and saved["mime_type"]=="text/html"
        assert "Gravity lab" in client.get("/api/v1/admin/resources/"+rid+"/preview").json()["html_document"]
        published=client.post("/api/v1/admin/resources/"+rid+"/publish",json={"expected_version":saved["version"]})
        assert published.status_code==200,published.text
        assert published.json()["published"]
