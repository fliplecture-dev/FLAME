"""Isolated Sheets adapter and rollback-only lifecycle; no live Google or provider calls."""
from copy import deepcopy
from uuid import UUID, uuid4
import pytest
from app.modules.bulk_transfer.google_sheets import collect, spreadsheet_id, recheck, MARKER
from app.modules.bulk_transfer.errors import BulkFailure
from test_curriculum_workspace import isolated_client


class FakeSheets:
    ident = "s" * 30
    def __init__(self):
        self.names={"Image":1,"Video":2,"Interactive":3}
        headers=["record_key","existing_id","title","external_url","description"]
        key=uuid4().hex
        self.grids={"Image":[headers.copy(),["","","Picture "+key,"https://media.example.test/"+key+".png","Diagram"]],
            "Video":[headers.copy(),["","","Video "+key,"https://media.example.test/"+key+".mp4","Motion"]],
            "Interactive":[headers.copy(),["","","Game "+key,"","Interactive draft"]]}
        self.markers={};self.offline=False;self.writes=0
    def read(self):return deepcopy((self.names,self.grids,self.markers))
    def mark(self,requests):
        result={}
        for x in requests:
            item={**x,"metadataId":len(self.markers)+100};loc=x["location"]["dimensionRange"]
            self.markers[(loc["sheetId"],loc["startIndex"]+1)]=item;result[x["metadataValue"]]=item["metadataId"]
        return result
    def call(self,method,suffix,**kwargs):
        if self.offline:raise BulkFailure(503,"Google unavailable for isolated test")
        if suffix==':batchUpdate':
            for request in kwargs['json']['requests']:
                if 'updateCells' in request:
                    action=request['updateCells'];tab=next(n for n,i in self.names.items() if i==action['start']['sheetId']);start=action['start']['rowIndex']
                    for index,row in enumerate(action['rows']):
                        while len(self.grids[tab])<=start+index:self.grids[tab].append([])
                        dest=self.grids[tab][start+index];col=action['start'].get('columnIndex',0)
                        while len(dest)<col+len(row['values']):dest.append('')
                        for n,x in enumerate(row['values']):dest[col+n]=next(iter(x['userEnteredValue'].values()))
                elif 'createDeveloperMetadata' in request:self.mark([request['createDeveloperMetadata']['developerMetadata']])
            return {}
        assert suffix=="/values:batchUpdateByDataFilter"
        self.writes+=1
        for data in kwargs["json"]["data"]:
            metadata_id=data["dataFilter"]["developerMetadataLookup"]["metadataId"]
            key=next(k for k,v in self.markers.items() if v["metadataId"]==metadata_id)
            tab=next(n for n,i in self.names.items() if i==key[0]);row=self.grids[tab][key[1]-1]
            for index,value in enumerate(data["values"][0]):
                while len(row)<=index:row.append('')
                if value is not None:row[index]=value
        return {"totalUpdatedCells":sum(sum(v is not None for v in d['values'][0]) for d in kwargs['json']['data'])}


def install(monkeypatch,fake):
    import app.modules.bulk_transfer.google_sheets as adapter
    import app.modules.bulk_transfer.google_router as router
    monkeypatch.setattr(adapter,"SheetsClient",lambda config,ident:fake)
    monkeypatch.setattr(router,"SheetsClient",lambda config,ident:fake)


def preview(client,fake):
    result=client.post('/api/v1/admin/bulk/google-sheets/preview',json={"spreadsheet":fake.ident,"connector":False})
    assert result.status_code==200,result.text
    return result.json()


def confirm(client,job):
    return client.post('/api/v1/admin/bulk/google-sheets/'+job['id']+'/commit',json={"preview_token":job['preview_token'],"confirm":True})


def test_url_columns_formulas_and_duplicate_keys():
    assert spreadsheet_id('https://docs.google.com/spreadsheets/d/'+'a'*30+'/edit#gid=0')=='a'*30
    with pytest.raises(BulkFailure):spreadsheet_id('https://evil.test/foo')
    fake=FakeSheets();fake.grids['Image'][0].append('titel')
    records,refs,issues,_=collect(*fake.read(),reserve=True)
    assert any(x['severity']=='error' and x['column']=='titel' and 'title' in x['message'] for x in issues)
    assert all(x['type']!='image' for x in records)
    fake=FakeSheets();fake.grids['Image'][1][0]='duplicate';fake.grids['Video'][1][0]='duplicate';fake.grids['Interactive'][1][2]='=SUM(A1)'
    records,refs,issues,_=collect(*fake.read(),reserve=True)
    assert 'Duplicate record_key' in refs[1]['errors'][0]
    assert 'Formulas' in refs[2]['errors'][0]
    del fake.names['Video']
    assert any(x['tab']=='Video' and x['severity']=='error' for x in collect(*fake.read(),reserve=True)[2])


def test_transport_reads_row_metadata_from_search():
    from app.modules.bulk_transfer.google_sheets import SheetsClient
    client=SheetsClient.__new__(SheetsClient);client.ident='s'*30
    calls=[]
    def call(method,suffix='',**kwargs):
        calls.append((method,suffix,kwargs))
        if not suffix:return {'sheets':[{'properties':{'sheetId':1,'title':'Image'}}]}
        if suffix=='/values:batchGet':return {'valueRanges':[{'values':[['record_key','existing_id','title'],['','','Fixture']]}]}
        return {'matchedDeveloperMetadata':[{'developerMetadata':{'metadataId':12,'metadataKey':MARKER,'metadataValue':'stable','location':{'dimensionRange':{'sheetId':1,'dimension':'ROWS','startIndex':1,'endIndex':2}}}}]}
    client.call=call
    names,grids,markers=client.read()
    assert markers[(1,2)]['metadataValue']=='stable'
    assert calls[-1][1]=='/developerMetadata:search' and calls[-1][2]['json']['dataFilters'][0]['developerMetadataLookup']['locationType']=='ROW'


def test_row_reorder_tracking_and_changed_content():
    fake=FakeSheets();records,refs,issues,requests=collect(*fake.read(),reserve=True);fake.mark(requests)
    source={'rows':refs}
    marker=fake.markers.pop((1,2));fake.grids['Image'].insert(1,['','','','','']);fake.markers[(1,3)]=marker
    assert len(recheck(fake,source))==3
    fake.grids['Image'][2][2]='Edited after preview'
    with pytest.raises(BulkFailure,match='source row changed'):recheck(fake,source)


def test_sync_lifecycle_failure_replay_and_mapping_preservation(monkeypatch):
    fake=FakeSheets();install(monkeypatch,fake)
    with isolated_client() as client:
        client.app.state.settings.resource_allowed_hosts='media.example.test'
        initial=preview(client,fake);job=initial['job'];assert job['counts']['add']==3
        assert len(job['sheet_issues'])==3  # optional column warnings
        fake.offline=True
        result=confirm(client,job);assert result.status_code==200,result.text
        saved=result.json();assert saved['status']=='completed' and saved['writeback']=='pending'
        ids=saved['result']['ids'];assert len(ids)==3
        assert confirm(client,job).json()['result']['ids']==ids
        fake.offline=False
        retry=client.post('/api/v1/admin/bulk/google-sheets/'+job['id']+'/writeback',json={})
        assert retry.json()['writeback']=='completed'
        assert fake.grids['Image'][1][0].startswith('gs-') and fake.grids['Image'][1][1] in ids.values()
        count=fake.writes;assert confirm(client,job).status_code==200 and fake.writes==count
        again=preview(client,fake)['job'];assert again['counts']['unchanged']==3
        assert confirm(client,again).json()['result']['added']==0
        # Existing publication and mappings survive a conflicting sheet edit.
        from test_bulk_transfer import path_fixture
        from app.modules.resources.models import Resource,ResourceTopic
        topic,leaf,subject=path_fixture(client)
        rid=fake.grids['Image'][1][1]
        with client.app.state.database.session() as db, db.begin():
            row=db.get(Resource,UUID(rid));row.published=True
            db.add(ResourceTopic(resource_id=row.id,topic_id=UUID(topic),subtopic_id=UUID(leaf)))
        fake.grids['Image'][1][2]='Different sheet title'
        third=preview(client,fake)['job'];assert 'title' in third['rows'][0]['conflicts']
        assert confirm(client,third).status_code==200
        retained=client.get('/api/v1/admin/resources/'+rid).json()
        assert retained['published'] and len(retained['paths'])==1 and retained['title']!='Different sheet title'
        # Removing a sheet row does not delete a resource.
        fake.grids['Image'].pop();fake.markers.pop((1,2))
        fourth=preview(client,fake)['job'];assert fourth['total']==2
        assert confirm(client,fourth).status_code==200
        assert client.get('/api/v1/admin/resources/'+rid).status_code==200


def test_changed_preview_cancel_and_invalid_confirmation(monkeypatch):
    fake=FakeSheets();install(monkeypatch,fake)
    with isolated_client() as client:
        client.app.state.settings.resource_allowed_hosts='media.example.test'
        job=preview(client,fake)['job'];fake.grids['Video'][1][4]='Changed description'
        rejected=confirm(client,job);assert rejected.status_code==409,rejected.text
        assert client.get('/api/v1/admin/bulk/imports/'+job['id']).json()['status']=='review'
        assert client.post('/api/v1/admin/bulk/imports/'+job['id']+'/cancel',json={}).status_code==200
        assert confirm(client,job).status_code==409
        fake.grids['Video'][1][3]=''
        job=preview(client,fake)['job'];assert job['counts']['invalid']==1
        assert confirm(client,job).status_code==422


def test_unconfigured_and_template():
    from io import BytesIO
    from openpyxl import load_workbook
    with isolated_client() as client:
        assert client.get('/api/v1/admin/bulk/google-sheets/configuration').json()['configured'] is False
        result=client.post('/api/v1/admin/bulk/google-sheets/preview',json={'spreadsheet':'s'*30})
        assert result.status_code==503
        book=load_workbook(BytesIO(client.get('/api/v1/admin/bulk/google-sheets/template').content))
        assert set(book.sheetnames)=={'Image','Video','Interactive','Guide'}
        assert list(book['Image'].values)[0][:2]==('record_key','existing_id')
